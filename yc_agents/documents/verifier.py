import json
import os
import shutil
import re
from statistics import median
from pathlib import Path

from docx import Document

from yc_agents.documents.ooxml import package_part_hashes, sha256_file, validate_docx_package
from yc_agents.documents.analyzer import _font_format, _paragraph_format, _role_for, DocxTemplateAnalyzer
from yc_agents.documents.builder import (
    _has_field,
    _heading_text_for_paragraph,
    _is_body_sample,
    _styled_heading_level,
)
from yc_agents.documents.outline import flatten_outline


ALLOWED_CHANGED_PARTS = {
    "word/document.xml",
    "word/settings.xml",
    "docProps/core.xml",
    "docProps/app.xml",
}


class DocxVerifier:
    def __init__(self, job_store, broker=None, vision_service=None):
        self.job_store = job_store
        self.broker = broker
        self.vision_service = vision_service

    def verify(self, job_id, version=None, mode="all"):
        job = self.job_store.get(job_id)
        revision = self.job_store.revision(job_id, version)
        version = int(revision["version"])
        docx_path = Path(revision["docx_path"])
        qa_dir = self.job_store.job_root(job_id) / "qa" / f"v{version:03d}"
        qa_dir.mkdir(parents=True, exist_ok=True)
        findings = []
        artifacts = []

        try:
            validate_docx_package(docx_path)
        except Exception as exc:
            findings.append(self._finding("blocking", "package", str(exc)))
        if sha256_file(Path(job["template"]["path"])) != job["template"]["sha256"]:
            findings.append(self._finding("blocking", "template", "模板快照哈希发生变化"))

        template_parts = package_part_hashes(job["template"]["path"])
        revision_parts = package_part_hashes(docx_path)
        allowed_changed_parts = ALLOWED_CHANGED_PARTS | set(revision.get("allowed_changed_parts") or [])
        for name, metadata in template_parts.items():
            if name in allowed_changed_parts or name.startswith("word/media/"):
                continue
            current = revision_parts.get(name)
            if current is None:
                findings.append(self._finding("blocking", name, "生成文档丢失模板部件"))
            elif current["sha256"] != metadata["sha256"]:
                findings.append(self._finding("blocking", name, "模板保留部件发生非预期变化"))

        spec = self._load_spec(job)
        try:
            generated = Document(docx_path)
            if len(generated.sections) != len(spec.get("sections", [])):
                findings.append(self._finding("warning", "sections", "生成文档分节数量与模板不同"))
            else:
                for expected, actual in zip(spec.get("sections", []), generated.sections):
                    if int(actual.page_width or 0) != int((expected.get("page_width") or {}).get("emu") or 0):
                        findings.append(self._finding("blocking", f"section-{expected['index']}", "页面宽度与模板不一致"))
                    if int(actual.page_height or 0) != int((expected.get("page_height") or {}).get("emu") or 0):
                        findings.append(self._finding("blocking", f"section-{expected['index']}", "页面高度与模板不一致"))
                    for key in [
                        "top_margin",
                        "bottom_margin",
                        "left_margin",
                        "right_margin",
                        "header_distance",
                        "footer_distance",
                    ]:
                        actual_value = int(getattr(actual, key, None) or 0)
                        expected_value = int((expected.get(key) or {}).get("emu") or 0)
                        if actual_value != expected_value:
                            findings.append(
                                self._finding(
                                    "blocking",
                                    f"section-{expected['index']}",
                                    f"页面设置 {key} 与模板不一致",
                                )
                            )
            findings.extend(self._representative_format_findings(spec, generated))
            findings.extend(self._table_geometry_findings(spec, generated, revision, job))
            findings.extend(self._structural_content_findings(job, spec, generated))
        except Exception as exc:
            findings.append(self._finding("blocking", "document", f"无法读取生成DOCX：{exc}"))
        findings.extend(self._font_findings(spec))

        pdf_path = qa_dir / "document.pdf"
        page_images = []
        render_result = None
        if mode in {"all", "render", "visual"}:
            if self.broker is None:
                findings.append(
                    self._finding(
                        "blocking",
                        "renderer",
                        "Word ExecutionBroker 未配置",
                        "环境问题：无法用 docx_edit 修复，请向用户说明",
                        category="environment",
                    )
                )
            else:
                render_result = self.broker.run("word_export_pdf", docx_path, pdf_path)
                (qa_dir / "render-result.json").write_text(
                    json.dumps(render_result, ensure_ascii=False, indent=2), encoding="utf-8"
                )
                if not render_result["ok"]:
                    findings.append(
                        self._finding(
                            "blocking",
                            "renderer",
                            render_result.get("stderr")
                            or render_result.get("stdout")
                            or "Word PDF导出失败",
                        )
                    )
                else:
                    artifacts.append(str(pdf_path))
                    page_images, geometry_findings = self._render_pdf_pages(pdf_path, qa_dir / "pages")
                    findings.extend(geometry_findings)
                    findings.extend(self._page_number_findings(pdf_path))
                    artifacts.extend(str(path) for path in page_images)

        if mode in {"all", "visual"} and page_images:
            vision = self.vision_service.inspect_pages(page_images, template_summary=self._template_summary(spec)) if self.vision_service else {"available": False, "findings": []}
            if not vision.get("available") and not any(
                item.get("severity") == "blocking" for item in vision.get("findings", [])
            ):
                vision.setdefault("findings", []).append(
                    self._finding(
                        "blocking",
                        "vision",
                        "视觉模型未配置，未执行逐页图片检查",
                        "环境问题：无法用 docx_edit 修复，请向用户说明",
                        category="environment",
                    )
                )
            for item in vision.get("findings", []):
                item.setdefault("category", "document")
            findings.extend(vision.get("findings", []))
            (qa_dir / "vision-result.json").write_text(
                json.dumps(vision, ensure_ascii=False, indent=2), encoding="utf-8"
            )

        blocking = [item for item in findings if item.get("severity") == "blocking"]
        passed = not blocking
        published_path = None
        if mode == "all" and passed:
            target_value = revision.get("pending_published_path") or revision.get("published_path")
            if not target_value:
                findings.append(self._finding("blocking", "publish", "修订版本缺少待发布路径"))
                passed = False
            else:
                target = Path(target_value)
                target.parent.mkdir(parents=True, exist_ok=True)
                if target.exists() and sha256_file(target) != sha256_file(docx_path):
                    findings.append(self._finding("blocking", "publish", "输出路径已存在不同内容，拒绝覆盖"))
                    passed = False
                else:
                    if not target.exists():
                        shutil.copyfile(docx_path, target)
                    published_path = str(target)
                    artifacts.append(published_path)
        blocking = [item for item in findings if item.get("severity") == "blocking"]
        passed = not blocking
        for item in findings:
            item.setdefault("category", "document")
        environment_blocked = any(
            item.get("category") == "environment" for item in blocking
        )
        report = {
            "passed": passed,
            "mode": mode,
            "job_id": job_id,
            "version": version,
            "findings": findings,
            "blocking_count": len(blocking),
            "environment_blocked": environment_blocked,
            "warning_count": len([item for item in findings if item.get("severity") == "warning"]),
            "pdf_path": str(pdf_path) if pdf_path.exists() else None,
            "page_images": [str(path) for path in page_images],
            "render": render_result,
            "published_path": published_path,
            "delivery_ready": bool(mode == "all" and passed and published_path),
        }
        report_path = qa_dir / ("qa-report.json" if mode == "all" else f"qa-report-{mode}.json")
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        artifacts.append(str(report_path))

        revisions = list(job.get("revisions") or [])
        for item in revisions:
            if int(item.get("version", -1)) == version:
                modes = dict(item.get("qa_modes") or {})
                modes[mode] = {"passed": passed, "qa_report_path": str(report_path)}
                item["qa_modes"] = modes
                if mode == "all":
                    item["qa_passed"] = passed
                    item["qa_report_path"] = str(report_path)
                    item["pdf_path"] = report["pdf_path"]
                    item["page_images"] = report["page_images"]
                    item["delivery_ready"] = report["delivery_ready"]
                    if report["delivery_ready"]:
                        item["published_path"] = published_path
                    elif not item.get("delivery_ready"):
                        pending = item.get("pending_published_path") or item.get("published_path")
                        item["pending_published_path"] = pending
                        item["published_path"] = None
                manifest_path = item.get("manifest_path")
                if manifest_path:
                    Path(manifest_path).write_text(
                        json.dumps(item, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
        next_status = job.get("status")
        if mode == "all":
            next_status = "waiting_revision" if passed else "failed"
        self.job_store.update(
            job_id,
            revisions=revisions,
            qa={**dict(job.get("qa") or {}), f"v{version:03d}:{mode}": report},
            status=next_status,
        )
        return {**report, "qa_report_path": str(report_path), "artifacts": artifacts}

    @staticmethod
    def _finding(severity, anchor, issue, suggested_action="", category="document"):
        return {
            "severity": severity,
            "page": None,
            "anchor": anchor,
            "issue": issue,
            "suggested_action": suggested_action,
            "category": category,
        }

    @staticmethod
    def _load_spec(job):
        path = job.get("template_spec_path")
        if not path or not Path(path).exists():
            return {"sections": []}
        return json.loads(Path(path).read_text(encoding="utf-8"))

    @staticmethod
    def _template_summary(spec):
        return {
            "sections": spec.get("sections", []),
            "format_clusters": spec.get("format_clusters", [])[:12],
            "tables": [
                {key: table.get(key) for key in ["element_id", "style", "rows", "columns", "grid_widths_dxa"]}
                for table in spec.get("tables", [])[:20]
            ],
        }

    def _render_pdf_pages(self, pdf_path, output_dir):
        try:
            import fitz
        except ImportError:
            return [], [
                self._finding(
                    "blocking",
                    "pdf-render",
                    "PDF逐页渲染需要PyMuPDF依赖",
                    "环境问题：无法用 docx_edit 修复，请向用户说明",
                    category="environment",
                )
            ]
        output_dir.mkdir(parents=True, exist_ok=True)
        document = fitz.open(pdf_path)
        paths = []
        findings = []
        for index, page in enumerate(document):
            page_number = index + 1
            for block in page.get_text("blocks"):
                x0, y0, x1, y1 = block[:4]
                if x0 < -1 or y0 < -1 or x1 > page.rect.width + 1 or y1 > page.rect.height + 1:
                    findings.append(
                        {
                            "severity": "blocking",
                            "page": page_number,
                            "anchor": str(block[4])[:80],
                            "issue": "PDF文本边界超出页面",
                            "suggested_action": "检查段落、表格或图片宽度",
                        }
                    )
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            path = output_dir / f"page-{page_number}.png"
            pixmap.save(path)
            paths.append(path)
        document.close()
        return paths, findings

    def _representative_format_findings(self, spec, generated):
        findings = []
        paragraphs = list(generated.paragraphs)
        nonempty = [index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip()]
        first_nonempty = nonempty[0] if nonempty else -1
        actual_by_role = {}
        for index, paragraph in enumerate(paragraphs):
            role = _role_for(paragraph, index, first_nonempty)
            if paragraph.text.strip() and role not in actual_by_role:
                actual_by_role[role] = paragraph
        expected_by_role = {}
        for element in spec.get("elements", []):
            role = element.get("role")
            if element.get("text", "").strip() and role not in expected_by_role:
                expected_by_role[role] = element
        for role in ["title", "heading_1", "heading_2", "body", "caption"]:
            expected = expected_by_role.get(role)
            paragraph = actual_by_role.get(role)
            if expected is None or paragraph is None:
                continue
            actual_paragraph = _paragraph_format(paragraph)
            expected_paragraph = expected.get("paragraph_format") or {}
            for key in ["alignment", "line_spacing_rule", "raw_spacing", "raw_indent"]:
                if actual_paragraph.get(key) != expected_paragraph.get(key):
                    findings.append(
                        self._finding(
                            "blocking",
                            role,
                            f"代表性{role}段落格式 {key} 与模板不一致",
                        )
                    )
            expected_runs = expected.get("runs") or []
            if expected_runs and paragraph.runs:
                expected_font = expected_runs[0].get("effective_font") or {}
                actual_font = _font_format(paragraph.runs[0], paragraph)
                for key in ["names", "size", "bold"]:
                    if actual_font.get(key) != expected_font.get(key):
                        findings.append(
                            self._finding(
                                "blocking",
                                role,
                                f"代表性{role}字符格式 {key} 与模板不一致",
                            )
                        )
        return findings

    def _structural_content_findings(self, job, spec, generated):
        findings = []
        paragraphs = list(generated.paragraphs)
        outline_sections = flatten_outline(job.get("outline") or {})
        expected_titles = {str(item.get("title") or "").strip() for item in outline_sections}

        title_groups = {}
        for section in outline_sections:
            title = str(section.get("title") or "").strip()
            title_groups.setdefault(title, []).append(section)
        for title, group in title_groups.items():
            level = int(group[0].get("level") or 1)
            matches = [
                paragraph
                for paragraph in paragraphs
                if paragraph.text.strip()
                == _heading_text_for_paragraph(title, paragraph, level)
            ]
            expected_count = len(group)
            if len(matches) != expected_count:
                findings.append(
                    self._finding(
                        "blocking",
                        title or group[0].get("id") or "outline",
                        f"提纲标题应在生成文档中出现 {expected_count} 次，实际为 {len(matches)} 次",
                        "标题缺失时检查章节标题是否被改写；多余时检查正文是否包含与标题完全相同的整行文本",
                    )
                )
                continue
            if expected_count == 1:
                actual_level = _styled_heading_level(matches[0])
                if actual_level != level:
                    findings.append(
                        self._finding(
                            "blocking",
                            title,
                            f"标题必须使用可导航的 Heading {level} 语义，实际层级为 {actual_level}",
                        )
                    )

        for paragraph in paragraphs:
            style_name = str(paragraph.style.name or "").strip().lower()
            text = paragraph.text.strip()
            if (style_name.startswith("toc") or style_name.startswith("目录")) and len(text) > 120:
                findings.append(
                    self._finding("blocking", text[:80], "目录样式段落包含长正文，正文被错误写入目录")
                )

        visible_text = "\n".join(
            [paragraph.text for paragraph in paragraphs]
            + [cell.text for table in generated.tables for row in table.rows for cell in row.cells]
        )
        if re.search(r"(?m)^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$", visible_text):
            findings.append(self._finding("blocking", "markdown-table", "文档中残留 Markdown 表格分隔线"))
        if re.search(r"错误\s*[!！]\s*(?:未定义书签|引用源未找到)|Error!\s*(?:Bookmark|Reference)", visible_text, re.IGNORECASE):
            findings.append(self._finding("blocking", "fields", "目录或交叉引用显示未定义书签错误"))

        root = generated._element
        bookmark_names = set(root.xpath(".//w:bookmarkStart/@w:name"))
        hyperlink_anchors = set(root.xpath(".//w:hyperlink/@w:anchor"))
        field_text = " ".join(root.xpath(".//w:instrText/text()"))
        referenced = set(hyperlink_anchors)
        referenced.update(
            match.group(1) or match.group(2)
            for match in re.finditer(
                r"PAGEREF\s+\"?([^\s\"\\]+)|HYPERLINK\s+\\l\s+\"([^\"]+)\"",
                field_text,
                re.IGNORECASE,
            )
        )
        missing_bookmarks = sorted(
            name for name in referenced if name and name not in bookmark_names and name != "_GoBack"
        )
        if missing_bookmarks:
            findings.append(
                self._finding(
                    "blocking",
                    "bookmarks",
                    f"目录或内部链接引用了不存在的书签：{', '.join(missing_bookmarks[:12])}",
                )
            )

        expected_sizes = []
        for element in spec.get("elements", []):
            if element.get("role") != "body":
                continue
            for run in element.get("runs") or []:
                size = ((run.get("effective_font") or {}).get("size") or {}).get("pt")
                if size:
                    expected_sizes.append(float(size))
        expected_size = median(expected_sizes) if expected_sizes else 12.0
        oversized_chars = 0
        oversized_anchors = []
        for paragraph in paragraphs:
            if not _is_body_sample(paragraph):
                continue
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                size = ((_font_format(run, paragraph).get("size") or {}).get("pt"))
                if size and float(size) > max(18.0, expected_size * 1.5):
                    oversized_chars += len(run.text)
                    if len(oversized_anchors) < 5:
                        oversized_anchors.append(run.text.strip()[:30])
        if oversized_chars >= 80:
            findings.append(
                self._finding(
                    "blocking",
                    "body-font-size",
                    f"检测到 {oversized_chars} 个正文字符字号异常偏大；模板正文基准约 {expected_size:g}pt",
                    f"检查：{'、'.join(oversized_anchors)}",
                )
            )

        findings.extend(self._legacy_content_findings(job, generated, expected_titles))
        assumption_sections = []
        sections_dir = self.job_store.job_root(job["id"]) / "content" / "sections"
        if sections_dir.exists():
            for path in sections_dir.glob("*.json"):
                try:
                    value = json.loads(path.read_text(encoding="utf-8"))
                except (OSError, json.JSONDecodeError):
                    continue
                if str(value.get("fact_status") or "").lower() == "assumption":
                    assumption_sections.append(str(value.get("title") or path.stem))
        if assumption_sections and not re.search(r"【假设】|\[假设\]|假设说明|暂按|测算假设", visible_text):
            findings.append(
                self._finding(
                    "blocking",
                    "assumptions",
                    f"含假设数据的章节未在正文显式标注：{', '.join(assumption_sections[:8])}",
                )
            )
        return findings

    def _legacy_content_findings(self, job, generated, expected_titles):
        template = Document(job["template"]["path"])
        template_paragraphs = list(template.paragraphs)
        generated_texts = {paragraph.text.strip() for paragraph in generated.paragraphs if paragraph.text.strip()}
        first_heading = next(
            (index for index, paragraph in enumerate(template_paragraphs) if _styled_heading_level(paragraph)),
            0,
        )
        stale = []
        for paragraph in template_paragraphs[first_heading:]:
            text = paragraph.text.strip()
            if (
                len(text) >= 24
                and text not in expected_titles
                and not _has_field(paragraph)
                and not str(paragraph.style.name or "").lower().startswith("toc")
                and text in generated_texts
            ):
                stale.append(text)
        if stale:
            return [
                self._finding(
                    "blocking",
                    stale[0][:80],
                    f"生成文档仍残留 {len(stale)} 段模板旧业务正文",
                    "删除或重写旧项目内容后重新验证",
                )
            ]

        requirements = dict(job.get("requirements") or {})
        brand_decision = str(
            requirements.get("brand_decision") or requirements.get("brand_info") or ""
        ).strip().lower()
        if brand_decision not in {"delete", "remove", "删除"}:
            return []
        company_pattern = re.compile(
            r"[\u3400-\u9fffA-Za-z0-9（）()·]{2,28}(?:集团有限公司|有限公司)"
        )
        template_text = "\n".join(
            [p.text for p in template_paragraphs]
            + [cell.text for table in template.tables for row in table.rows for cell in row.cells]
        )
        old_companies = set(company_pattern.findall(template_text))
        generated_text = "\n".join(
            [p.text for p in generated.paragraphs]
            + [cell.text for table in generated.tables for row in table.rows for cell in row.cells]
        )
        allowed = " ".join(
            str(requirements.get(key) or job.get(key) or "")
            for key in ("company_name", "title", "topic")
        )
        residual = sorted(
            name for name in old_companies if name in generated_text and name not in allowed
        )
        return [
            self._finding(
                "blocking",
                "cover-brand",
                f"用户要求删除旧品牌，但仍残留：{', '.join(residual)}",
            )
        ] if residual else []

    def _page_number_findings(self, pdf_path):
        try:
            import fitz
        except ImportError:
            return []
        document = fitz.open(pdf_path)
        detected = []
        try:
            for page_index, page in enumerate(document, start=1):
                candidates = []
                for block in page.get_text("blocks"):
                    if block[1] < page.rect.height * 0.9:
                        continue
                    for line in str(block[4] or "").splitlines():
                        match = re.fullmatch(r"\s*(?:第\s*)?(\d{1,4})\s*(?:页)?\s*", line)
                        if match:
                            candidates.append(int(match.group(1)))
                if candidates:
                    detected.append((page_index, candidates[-1]))
        finally:
            document.close()
        for (previous_page, previous_number), (page, number) in zip(detected, detected[1:]):
            if page == previous_page + 1 and number != previous_number + 1:
                return [
                    self._finding(
                        "blocking",
                        f"page-{page}",
                        f"页码不连续或发生重置：PDF第{previous_page}页显示{previous_number}，第{page}页显示{number}",
                    )
                ]
        return []

    def _table_geometry_findings(self, spec, generated, revision, job):
        if revision.get("base") != "template":
            return []
        contract = {}
        path = job.get("template_contract_path")
        if path and Path(path).exists():
            contract = json.loads(Path(path).read_text(encoding="utf-8"))
        expected_tables = list(spec.get("tables") or [])
        action_by_id = {
            str(item.get("element_id")): str(item.get("action") or "confirm")
            for item in list(contract.get("tables") or [])
            if isinstance(item, dict) and item.get("element_id")
        }
        default_action = str((contract.get("defaults") or {}).get("tables") or "confirm")
        findings = []
        generated_specs = [
            DocxTemplateAnalyzer._table_spec(table, index)
            for index, table in enumerate(generated.tables)
        ]
        cursor = 0
        for index, expected in enumerate(expected_tables):
            element_id = str(expected.get("element_id") or f"body.tbl{index:04d}")
            action = action_by_id.get(element_id, default_action)
            if action == "delete":
                continue
            geometry_keys = ["columns", "grid_widths_dxa", "width", "indent", "cell_margins"]
            match_index = next(
                (
                    candidate
                    for candidate in range(cursor, len(generated_specs))
                    if all(
                        generated_specs[candidate].get(key) == expected.get(key)
                        for key in geometry_keys
                    )
                ),
                None,
            )
            if match_index is None:
                findings.append(
                    self._finding("blocking", element_id, "模板契约要求保留或复用的表格缺失")
                )
                continue
            actual = generated_specs[match_index]
            cursor = match_index + 1
            for key in geometry_keys:
                if actual.get(key) != expected.get(key):
                    findings.append(
                        self._finding(
                            "blocking",
                            element_id,
                            f"表格几何 {key} 与模板不一致",
                        )
                    )
        return findings

    def _font_findings(self, spec):
        installed = self._installed_font_names()
        if installed is None:
            return []
        requested = set()
        for element in [*spec.get("elements", []), *spec.get("headers", []), *spec.get("footers", [])]:
            for run in element.get("runs") or []:
                font = run.get("effective_font") or {}
                if font.get("name"):
                    requested.add(str(font["name"]))
                requested.update(
                    str(value)
                    for key, value in (font.get("names") or {}).items()
                    if not str(key).lower().endswith("theme")
                )
        missing = sorted(
            name
            for name in requested
            if name and "theme" not in name.lower() and not self._font_matches(name, installed)
        )
        return [
            self._finding(
                "warning",
                "fonts",
                f"模板字体未安装，Word 可能使用替代字体，无法声明字体级高保真：{', '.join(missing)}",
                "环境问题：安装字体后重新验证，或在交付说明中告知用户",
                category="environment",
            )
        ] if missing else []

    @staticmethod
    def _font_matches(requested, installed):
        normalized = re.sub(r"\s+", "", requested).casefold()
        aliases = {
            "黑体": {"simhei"},
            "宋体": {"simsun", "nsimsun"},
            "新宋体": {"nsimsun"},
            "仿宋": {"fangsong"},
            "楷体": {"kaiti"},
            "楷体_gb2312": {"kaiti", "stkaiti", "simkai"},
            "微软雅黑": {"microsoftyahei"},
            "等线": {"dengxian"},
        }
        candidates = {normalized, *aliases.get(normalized, set())}
        return any(
            candidate == item or candidate in item or item in candidate
            for candidate in candidates
            for item in installed
        )

    @staticmethod
    def _installed_font_names():
        if os.name != "nt":
            return None
        try:
            import winreg
        except ImportError:
            return None
        names = set()
        locations = [
            (winreg.HKEY_LOCAL_MACHINE, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"),
            (winreg.HKEY_CURRENT_USER, r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"),
        ]
        for hive, key_path in locations:
            try:
                with winreg.OpenKey(hive, key_path) as key:
                    index = 0
                    while True:
                        try:
                            name, _value, _kind = winreg.EnumValue(key, index)
                        except OSError:
                            break
                        cleaned = re.sub(r"\s*\([^)]*\)\s*$", "", name)
                        for alias in re.split(r"\s*[&,]\s*", cleaned):
                            if alias:
                                names.add(re.sub(r"\s+", "", alias).casefold())
                        index += 1
            except OSError:
                continue
        return names
