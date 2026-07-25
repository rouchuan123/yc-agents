import json
import re
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree
from pypdf import PdfReader

from yc_agents.rag.chunker import DocumentChunker
from yc_agents.rag.keyword_index import KeywordIndex, keyword_tokens
from yc_agents.documents.ooxml import sha256_file
from yc_agents.tools.readable_files import (
    is_blocked_readable_file,
    is_readable_document_file,
    is_readable_text_file,
    is_readable_workspace_file,
)


W_NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
EXCLUDED_DIRS = {".git", ".ycore", "node_modules", ".venv", "venv", "dist", "build", "__pycache__"}


class DocumentSourceService:
    def __init__(self, workspace_root, job_store, chunk_size=1200, chunk_overlap=150):
        self.workspace_root = Path(workspace_root).resolve()
        self.job_store = job_store
        self.chunker = DocumentChunker(chunk_size=chunk_size, overlap=chunk_overlap)

    def discover(self, job_id, query="", max_results=20, include_template=False):
        job = self.job_store.get(job_id)
        query_terms = set(keyword_tokens(query))
        candidates = []
        for path in self.workspace_root.rglob("*"):
            if not path.is_file() or set(path.relative_to(self.workspace_root).parts) & EXCLUDED_DIRS:
                continue
            if path.name.startswith("~$"):
                continue
            if is_blocked_readable_file(path) or not is_readable_workspace_file(path):
                continue
            if not include_template and self._is_template_source(path, job):
                continue
            relative = str(path.relative_to(self.workspace_root)).replace("\\", "/")
            terms = set(keyword_tokens(relative))
            score = len(query_terms & terms) / max(1, len(query_terms)) if query_terms else 0.0
            candidates.append(
                {
                    "id": f"src_{len(candidates) + 1:04d}",
                    "path": relative,
                    "name": path.name,
                    "suffix": path.suffix.lower(),
                    "bytes": path.stat().st_size,
                    "score": round(score, 4),
                }
            )
        candidates.sort(key=lambda item: (-item["score"], item["path"]))
        selected = candidates[: max(1, min(int(max_results), 100))]
        self.job_store.update(
            job_id,
            source_candidates=selected,
            status="waiting_source_confirmation",
        )
        return {"candidates": selected, "count": len(selected), "requires_confirmation": True}

    @staticmethod
    def _is_template_source(path, job):
        template = job.get("template") or {}
        template_path = Path(template.get("path") or "")
        try:
            if template_path and path.resolve() == template_path.resolve():
                return True
        except (OSError, RuntimeError):
            pass
        if path.name.casefold() != str(template.get("name") or "").casefold():
            return False
        try:
            return sha256_file(path) == template.get("sha256")
        except OSError:
            return False

    def confirm(self, job_id, source_ids):
        job = self.job_store.get(job_id)
        requested = set(source_ids or [])
        selected = [item for item in job.get("source_candidates", []) if item.get("id") in requested]
        missing = requested - {item["id"] for item in selected}
        if missing:
            raise ValueError(f"Unknown source candidate ids: {sorted(missing)}")
        self.job_store.update(
            job_id,
            confirmed_sources=selected,
            source_confirmation_at=self._now_iso(),
        )
        return {"confirmed_sources": selected, "count": len(selected)}

    def ingest(self, job_id):
        job = self.job_store.get(job_id)
        confirmed = list(job.get("confirmed_sources") or [])
        output = []
        chunks = []
        for source in confirmed:
            path = self._resolve_workspace_path(source["path"])
            text = self._read_source(path)
            source_chunks = self.chunker.chunk_text(
                text,
                source=source["id"],
                metadata={"path": source["path"], "name": source["name"]},
            )
            chunks.extend(
                {
                    "source": chunk.source,
                    "chunk_id": chunk.chunk_id,
                    "text": chunk.text,
                    "metadata": chunk.metadata,
                }
                for chunk in source_chunks
            )
            output.append({**source, "characters": len(text), "chunks": len(source_chunks)})
        root = self.job_store.job_root(job_id) / "sources"
        root.mkdir(parents=True, exist_ok=True)
        (root / "chunks.json").write_text(json.dumps(chunks, ensure_ascii=False, indent=2), encoding="utf-8")
        (root / "ingested.json").write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
        return {"sources": output, "chunks": len(chunks)}

    def search(self, job_id, query, top_k=5):
        chunks_path = self.job_store.job_root(job_id) / "sources" / "chunks.json"
        if not chunks_path.exists():
            self.ingest(job_id)
        chunks = json.loads(chunks_path.read_text(encoding="utf-8"))
        index = KeywordIndex()
        index.add_chunks("document-job", [item["text"] for item in chunks])
        results = index.search(query, top_k=max(1, min(int(top_k), 20)))
        for result in results:
            chunk_id = int(result["chunk_id"])
            if 0 <= chunk_id < len(chunks):
                source_chunk = chunks[chunk_id]
                result["source"] = source_chunk["source"]
                result["metadata"] = source_chunk.get("metadata", {})
        return {"query": query, "results": results, "count": len(results)}

    def record_web(self, job_id, source):
        root = self.job_store.job_root(job_id) / "sources"
        root.mkdir(parents=True, exist_ok=True)
        path = root / "web.json"
        items = json.loads(path.read_text(encoding="utf-8")) if path.exists() else []
        record = dict(source or {})
        record.setdefault("id", f"web_{len(items) + 1:04d}")
        items.append(record)
        path.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
        return record

    def _resolve_workspace_path(self, relative_path):
        candidate = (self.workspace_root / relative_path).resolve()
        if candidate != self.workspace_root and self.workspace_root not in candidate.parents:
            raise PermissionError(f"Source path escapes active workspace: {relative_path}")
        if not candidate.exists() or not candidate.is_file():
            raise FileNotFoundError(f"Confirmed source no longer exists: {relative_path}")
        return candidate

    def _read_source(self, path):
        if is_readable_text_file(path):
            return path.read_text(encoding="utf-8", errors="replace")
        if is_readable_document_file(path):
            if path.suffix.lower() == ".docx":
                return self._read_docx(path)
            if path.suffix.lower() == ".pdf":
                return "\n".join((page.extract_text() or "").strip() for page in PdfReader(path).pages)
        raise ValueError(f"Unsupported confirmed source: {path.name}")

    @staticmethod
    def _read_docx(path):
        document = Document(path)
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            blocks.extend("\t".join(cell.text.strip() for cell in row.cells) for row in table.rows)
        for section in document.sections:
            blocks.extend(paragraph.text for paragraph in section.header.paragraphs if paragraph.text.strip())
            blocks.extend(paragraph.text for paragraph in section.footer.paragraphs if paragraph.text.strip())
        with zipfile.ZipFile(path) as package:
            for name in package.namelist():
                if not name.endswith(".xml"):
                    continue
                try:
                    root = etree.fromstring(package.read(name))
                except etree.XMLSyntaxError:
                    continue
                for textbox in root.xpath("//w:txbxContent", namespaces=W_NS):
                    text = "".join(textbox.xpath(".//w:t/text()", namespaces=W_NS)).strip()
                    if text:
                        blocks.append(text)
        return "\n".join(blocks)

    @staticmethod
    def _now_iso():
        from datetime import datetime

        return datetime.now().isoformat(timespec="seconds")
