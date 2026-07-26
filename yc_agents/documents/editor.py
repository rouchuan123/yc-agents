import json
import re
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt

from yc_agents.documents.builder import (
    _content_paragraphs,
    _heading_level,
    _is_heading,
    _replace_paragraph_text,
    _request_field_update,
    _table_set_text,
)
from yc_agents.documents.ooxml import package_part_hashes, preserve_package_parts, sha256_file


def _now_iso():
    return datetime.now().isoformat(timespec="seconds")


class DocxEditor:
    def __init__(self, workspace_root, job_store):
        self.workspace_root = Path(workspace_root).resolve()
        self.job_store = job_store

    def edit(self, job_id, base_revision, operations, output_name=""):
        job = self.job_store.get(job_id)
        if int(job.get("current_revision") or 0) != int(base_revision):
            raise ValueError(
                f"Revision conflict: current is v{int(job.get('current_revision') or 0):03d}, "
                f"requested base is v{int(base_revision):03d}"
            )
        if not list(operations or []):
            raise ValueError(
                "docx_edit requires at least one operation; an empty edit would create a "
                "byte-identical revision."
            )
        base = self.job_store.revision(job_id, base_revision)
        version = self.job_store.next_version(job_id)
        revision_dir = self.job_store.job_root(job_id) / "revisions" / f"v{version:03d}"
        revision_dir.mkdir(parents=True, exist_ok=False)
        try:
            return self._edit_into(
                job, job_id, base, base_revision, version, revision_dir, operations, output_name
            )
        except BaseException:
            # Never leave a half-built revision directory behind; orphan dirs block later versions.
            shutil.rmtree(revision_dir, ignore_errors=True)
            raise

    def _edit_into(self, job, job_id, base, base_revision, version, revision_dir, operations, output_name):
        internal_docx = revision_dir / "document.docx"
        shutil.copyfile(base["docx_path"], internal_docx)
        document = Document(internal_docx)

        results = []
        for operation in list(operations or []):
            results.append(self._apply(document, operation))
        _request_field_update(document)
        document.save(internal_docx)
        changed_parts = {"word/document.xml", "word/settings.xml"}
        changed_parts.update(
            str(result["package_part"]).lstrip("/")
            for result in results
            if result.get("package_part")
        )
        preserve_package_parts(base["docx_path"], internal_docx, changed_parts)

        output_dir = self.workspace_root / "outputs" / job["slug"]
        output_dir.mkdir(parents=True, exist_ok=True)
        filename = Path(output_name or f"{job['slug']}-v{version:03d}.docx").name
        if not filename.lower().endswith(".docx"):
            filename += ".docx"
        published = output_dir / filename
        if published.exists():
            raise FileExistsError(f"Published document version already exists: {published}")

        manifest = {
            "version": version,
            "base": int(base_revision),
            "operations": list(operations or []),
            "operation_results": results,
            "template_sha256": job["template"]["sha256"],
            "docx_path": str(internal_docx),
            "published_path": None,
            "pending_published_path": str(published),
            "docx_sha256": sha256_file(internal_docx),
            "package_parts": package_part_hashes(internal_docx),
            "allowed_changed_parts": list(base.get("allowed_changed_parts") or changed_parts),
            "created_at": _now_iso(),
            "qa_passed": False,
            "delivery_ready": False,
        }
        manifest_path = revision_dir / "artifact-manifest.json"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        manifest["manifest_path"] = str(manifest_path)
        self.job_store.add_revision(job_id, manifest)
        self.job_store.update(job_id, status="verifying")
        return {
            "ok": True,
            "job_id": job_id,
            "version": version,
            "docx_path": str(internal_docx),
            "published_path": None,
            "pending_published_path": str(published),
            "delivery_ready": False,
            "operations": results,
            "artifacts": [str(manifest_path)],
        }

    _OPERATION_ALIASES = {
        "replace": "replace_text",
        "insert_paragraph": "insert",
        "add_paragraph": "insert",
        "insert_text": "insert",
        "delete_paragraph": "delete",
        "delete_table": "delete",
        "remove": "delete",
        "edit_table": "update_table",
        "replace_table": "update_table",
        "set_format": "set_style",
        "update_style": "set_style",
    }

    def _apply(self, document, operation):
        if not isinstance(operation, dict):
            raise ValueError("Each edit operation must be an object")
        name = str(operation.get("operation") or "").strip().lower()
        name = self._OPERATION_ALIASES.get(name, name)
        if name == "replace_text":
            return self._replace_text(document, operation)
        if name == "replace_section":
            return self._replace_section(document, operation)
        if name == "insert":
            return self._insert(document, operation)
        if name == "delete":
            return self._delete(document, operation)
        if name == "move":
            return self._move(document, operation)
        if name == "update_table":
            return self._update_table(document, operation)
        if name == "delete_table_column":
            return self._delete_table_column(document, operation)
        if name == "set_style":
            return self._set_style(document, operation)
        if name == "replace_image":
            return self._replace_image(document, operation)
        raise ValueError(
            f"Unsupported DOCX edit operation: {name}. Supported operations: replace_text, "
            "replace_section, insert, delete, move, update_table, delete_table_column, "
            "set_style, replace_image."
        )

    @staticmethod
    def _all_paragraphs(document):
        paragraphs = list(document.paragraphs)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.extend(cell.paragraphs)
        for section in document.sections:
            # Accessing an unreferenced header/footer through python-docx
            # materializes a new relationship. A read-only target scan must
            # not mutate the package or introduce dangling rIds.
            sect_pr = section._sectPr
            if sect_pr.findall(qn("w:headerReference")):
                paragraphs.extend(section.header.paragraphs)
            if sect_pr.findall(qn("w:footerReference")):
                paragraphs.extend(section.footer.paragraphs)
        return paragraphs

    def _find_paragraph(self, document, target):
        target = str(target or "")
        match = re.fullmatch(r"body\.p(\d+)", target)
        if match:
            index = int(match.group(1))
            if index < len(document.paragraphs):
                return document.paragraphs[index]
        matches = [paragraph for paragraph in self._all_paragraphs(document) if target and target in paragraph.text]
        if len(matches) != 1:
            raise ValueError(f"Paragraph target must resolve exactly once, found {len(matches)}: {target}")
        return matches[0]

    @staticmethod
    def _is_table_id(target):
        return bool(re.search(r"(?:^|\.)(?:tbl|table)[-_\.]?\d+$", str(target or "")))

    def _find_table(self, document, target):
        target = str(target or "")
        match = re.search(r"(?:tbl|table)[-_\.]?(\d+)$", target)
        if match:
            index = int(match.group(1))
            if 0 <= index < len(document.tables):
                return document.tables[index]
            available = ", ".join(f"body.tbl{i:04d}" for i in range(len(document.tables)))
            raise ValueError(
                f"Table target out of range: {target}. Table ids are zero-based; "
                f"available: {available or 'none'}"
            )
        matches = [table for table in document.tables if any(target in cell.text for row in table.rows for cell in row.cells)]
        if len(matches) != 1:
            raise ValueError(
                f"Table target must resolve exactly once, found {len(matches)}: {target}. "
                "Use the element id (e.g. body.tbl0000) or unique cell text."
            )
        return matches[0]

    def _replace_text(self, document, operation):
        old = str(operation.get("old_text") or operation.get("target") or "")
        new = str(operation.get("new_text") or "")
        if not old:
            raise ValueError("replace_text requires old_text or target")
        paragraphs = self._all_paragraphs(document)
        occurrences = sum(paragraph.text.count(old) for paragraph in paragraphs)
        expected = int(operation.get("expected_replacements", 1))
        if occurrences != expected:
            raise ValueError(
                f"replace_text expected {expected} occurrence(s) of the text, found {occurrences}. "
                "Nothing was changed; adjust expected_replacements or make old_text more specific."
            )
        touched = 0
        for paragraph in paragraphs:
            if old in paragraph.text:
                _replace_paragraph_text(paragraph, paragraph.text.replace(old, new))
                touched += 1
        return {"operation": "replace_text", "replacements": occurrences, "paragraphs": touched}

    def _replace_section(self, document, operation):
        heading = self._find_paragraph(document, operation.get("target"))
        body_paragraphs = list(document.paragraphs)
        heading_index = next(
            (index for index, item in enumerate(body_paragraphs) if item._p is heading._p),
            None,
        )
        if heading_index is None:
            raise ValueError("replace_section target must be a body paragraph")
        source_level = _heading_level(heading) or 1
        next_heading = next(
            (
                item
                for item in body_paragraphs[heading_index + 1 :]
                if _is_heading(item) and (_heading_level(item) or 1) <= source_level
            ),
            None,
        )
        children = list(heading._p.getparent())
        start = children.index(heading._p)
        end = children.index(next_heading._p) if next_heading is not None else len(children)
        body = [item for item in children[start + 1 : end] if item.tag == qn("w:p")]
        sample = body[0] if body else deepcopy(heading._p)
        for item in body:
            item.getparent().remove(item)
        anchor = heading._p
        for text in _content_paragraphs(operation.get("content")):
            paragraph_xml = deepcopy(sample)
            anchor.addnext(paragraph_xml)
            paragraph = document.paragraphs[0].__class__(paragraph_xml, heading._parent)
            _replace_paragraph_text(paragraph, text)
            anchor = paragraph_xml
        if operation.get("title"):
            _replace_paragraph_text(heading, operation["title"])
        return {"operation": "replace_section", "target": operation.get("target")}

    def _insert(self, document, operation):
        target = self._find_paragraph(document, operation.get("target"))
        xml = deepcopy(target._p)
        target._p.addnext(xml)
        paragraph = target.__class__(xml, target._parent)
        _replace_paragraph_text(paragraph, operation.get("content", ""))
        return {"operation": "insert", "target": operation.get("target")}

    def _delete(self, document, operation):
        target = str(operation.get("target") or "")
        kind = str(operation.get("type") or "").strip().lower()
        if kind == "table" or (not kind and self._is_table_id(target)):
            table = self._find_table(document, target)
            table._tbl.getparent().remove(table._tbl)
        else:
            paragraph = self._find_paragraph(document, target)
            paragraph._p.getparent().remove(paragraph._p)
        return {"operation": "delete", "target": target}

    def _move(self, document, operation):
        source = self._find_paragraph(document, operation.get("target"))
        destination = self._find_paragraph(document, operation.get("before"))
        source._p.getparent().remove(source._p)
        destination._p.addprevious(source._p)
        return {"operation": "move", "target": operation.get("target"), "before": operation.get("before")}

    def _update_table(self, document, operation):
        table = self._find_table(document, operation.get("target"))
        values = list(operation.get("rows") or [])
        while len(table.rows) < len(values):
            table._tbl.append(deepcopy(table.rows[-1]._tr))
        while len(table.rows) > len(values):
            row = table.rows[-1]._tr
            row.getparent().remove(row)
        for row_index, row_values in enumerate(values):
            for column_index, value in enumerate(row_values):
                if column_index < len(table.rows[row_index].cells):
                    _table_set_text(table.rows[row_index].cells[column_index], value)
        return {"operation": "update_table", "rows": len(values)}

    def _delete_table_column(self, document, operation):
        table = self._find_table(document, operation.get("target"))
        column = operation.get("column")
        if isinstance(column, str) and not column.isdigit():
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            if column not in headers:
                raise ValueError(f"Table column not found: {column}")
            column_index = headers.index(column)
        else:
            column_index = int(column)
        grid = table._tbl.tblGrid
        grid_columns = grid.findall(qn("w:gridCol")) if grid is not None else []
        if 0 <= column_index < len(grid_columns):
            grid.remove(grid_columns[column_index])
        for row in table._tbl.tr_lst:
            cells = row.tc_lst
            if not 0 <= column_index < len(cells):
                raise ValueError(f"Table column index out of range: {column_index}")
            row.remove(cells[column_index])
        return {"operation": "delete_table_column", "column": column_index}

    def _set_style(self, document, operation):
        paragraph = self._find_paragraph(document, operation.get("target"))
        style = dict(operation.get("style") or {})
        if style.get("style_id"):
            paragraph.style = style["style_id"]
        if style.get("alignment") is not None:
            paragraph.alignment = int(style["alignment"])
        if style.get("line_spacing_pt") is not None:
            paragraph.paragraph_format.line_spacing = Pt(float(style["line_spacing_pt"]))
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        if style.get("first_line_indent_pt") is not None:
            paragraph.paragraph_format.first_line_indent = Pt(float(style["first_line_indent_pt"]))
        for run in paragraph.runs:
            if style.get("font_name"):
                run.font.name = style["font_name"]
                run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), style["font_name"])
            if style.get("font_size_pt") is not None:
                run.font.size = Pt(float(style["font_size_pt"]))
            if style.get("bold") is not None:
                run.bold = bool(style["bold"])
        return {"operation": "set_style", "target": operation.get("target")}

    def _replace_image(self, document, operation):
        image_path = Path(str(operation.get("image_path") or "")).resolve()
        if not image_path.exists() or not image_path.is_file():
            raise FileNotFoundError(f"Replacement image not found: {image_path}")
        image_parts = [part for part in document.part.package.parts if part.partname.startswith("/word/media/")]
        target = str(operation.get("target") or "image-0")
        match = re.search(r"(\d+)$", target)
        index = int(match.group(1)) if match else 0
        if not 0 <= index < len(image_parts):
            raise ValueError(f"Image target out of range: {target}")
        image_parts[index]._blob = image_path.read_bytes()
        return {
            "operation": "replace_image",
            "target": target,
            "bytes": image_path.stat().st_size,
            "package_part": str(image_parts[index].partname).lstrip("/"),
        }
