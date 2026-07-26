import json
import re
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from yc_agents.documents.content import (
    is_literature_review_job,
)
from yc_agents.documents.contract import normalize_template_contract
from yc_agents.documents.ooxml import package_part_hashes, preserve_package_parts, sha256_file


def _now_iso():
    return datetime.now().isoformat(timespec="seconds")


def _clean_text(value):
    value = str(value or "").strip()
    return re.sub(r"^(?:#{1,6}\s+|[-*]\s+)", "", value)


def _content_paragraphs(content):
    content = str(content or "").replace("\r\n", "\n")
    blocks = [block.strip() for block in re.split(r"\n\s*\n", content) if block.strip()]
    if len(blocks) <= 1:
        blocks = [line.strip() for line in content.splitlines() if line.strip()]
    return [_clean_text(block) for block in blocks]


def _replace_paragraph_text(paragraph, text):
    text = str(text or "")
    text_nodes = paragraph._p.xpath(".//w:t")
    if text_nodes:
        text_nodes[0].text = text
        if text.startswith(" ") or text.endswith(" "):
            text_nodes[0].set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        for node in text_nodes[1:]:
            node.text = ""
        return
    run = paragraph.add_run(text)
    if paragraph.runs and len(paragraph.runs) > 1:
        source_rpr = paragraph.runs[0]._r.rPr
        if source_rpr is not None:
            run._r.insert(0, deepcopy(source_rpr))


def _styled_heading_level(paragraph):
    text = paragraph.text.strip()
    style_name = (paragraph.style.name or "").lower()
    if not text:
        return None
    if style_name.startswith("toc") or style_name.startswith("目录"):
        return None
    style_match = re.search(r"(?:heading|标题)\s*([1-9])", style_name)
    if style_match:
        return int(style_match.group(1))
    outline_level = paragraph._p.xpath("./w:pPr/w:outlineLvl/@w:val")
    if outline_level:
        try:
            return int(outline_level[0]) + 1
        except (TypeError, ValueError):
            pass
    return None


def _heading_level(paragraph):
    styled = _styled_heading_level(paragraph)
    if styled is not None:
        return styled
    text = paragraph.text.strip()
    style_name = (paragraph.style.name or "").lower()
    if not text or style_name.startswith("toc") or style_name.startswith("目录"):
        return None
    if _has_field(paragraph):
        return None
    if re.match(
        r"^(?:第[一二三四五六七八九十百0-9]+[章节篇部](?:\s|[、：:])|[一二三四五六七八九十百]+、)",
        text,
    ):
        return 1
    numeric = re.match(r"^(\d+(?:\.\d+){0,3})[、.\s]", text)
    if numeric:
        return numeric.group(1).count(".") + 1
    if "heading" in style_name or "标题" in style_name:
        return 1
    return None


def _is_heading(paragraph):
    return _heading_level(paragraph) is not None


def _has_field(paragraph):
    return bool(paragraph._p.xpath(".//w:instrText | .//w:fldChar"))


def _is_body_sample(paragraph):
    style_name = (paragraph.style.name or "").strip().lower()
    text = paragraph.text.strip()
    direct_sizes = [run.font.size.pt for run in paragraph.runs if run.font.size is not None]
    return bool(
        text
        and not _is_heading(paragraph)
        and not _has_field(paragraph)
        and not style_name.startswith("toc")
        and style_name not in {"title", "subtitle", "题名", "副标题"}
        and (not direct_sizes or max(direct_sizes) <= 18)
        and not re.match(r"^(?:图|表)\s*\d+", text)
    )


def _markdown_row(line):
    return [cell.strip() for cell in str(line).strip().strip("|").split("|")]


def _is_markdown_separator(line):
    cells = _markdown_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def _content_blocks(content):
    lines = str(content or "").replace("\r\n", "\n").splitlines()
    blocks = []
    prose = []

    def flush_prose():
        if prose:
            for text in _content_paragraphs("\n".join(prose)):
                blocks.append(("paragraph", text))
            prose.clear()

    index = 0
    while index < len(lines):
        line = lines[index]
        if (
            line.strip().startswith("|")
            and index + 1 < len(lines)
            and _is_markdown_separator(lines[index + 1])
        ):
            flush_prose()
            headers = _markdown_row(line)
            index += 2
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append(_markdown_row(lines[index]))
                index += 1
            blocks.append(("table", {"headers": headers, "rows": rows}))
            continue
        if line.strip():
            prose.append(line)
        else:
            flush_prose()
        index += 1
    flush_prose()
    return blocks


def _table_set_text(cell, text):
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _replace_paragraph_text(paragraph, text)
    for extra in list(cell.paragraphs[1:]):
        extra._element.getparent().remove(extra._element)


def _request_field_update(document):
    settings = document.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def _continue_page_numbering(document):
    for section in list(document.sections)[1:]:
        page_number = section._sectPr.find(qn("w:pgNumType"))
        if page_number is None:
            continue
        page_number.attrib.pop(qn("w:start"), None)
        if not page_number.attrib and len(page_number) == 0:
            section._sectPr.remove(page_number)


def _normalize_floating_footer_page_fields(document):
    alternate_content_tag = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
    changed_parts = set()
    seen_parts = set()
    for section in document.sections:
        footer = section.footer
        part_name = str(footer.part.partname).lstrip("/")
        if part_name in seen_parts:
            continue
        seen_parts.add(part_name)
        for paragraph in footer.paragraphs:
            alternates = list(paragraph._p.iter(alternate_content_tag))
            page_alternates = [
                node
                for node in alternates
                if any(
                    "PAGE" in str(instruction.text or "").upper()
                    for instruction in node.iter(qn("w:instrText"))
                )
            ]
            if not page_alternates:
                continue
            for node in page_alternates:
                parent = node.getparent()
                if parent is not None:
                    parent.remove(node)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            begin = paragraph.add_run()._r
            begin_char = OxmlElement("w:fldChar")
            begin_char.set(qn("w:fldCharType"), "begin")
            begin.append(begin_char)
            instruction = paragraph.add_run()._r
            instruction_text = OxmlElement("w:instrText")
            instruction_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            instruction_text.text = " PAGE  \\* MERGEFORMAT "
            instruction.append(instruction_text)
            separate = paragraph.add_run()._r
            separate_char = OxmlElement("w:fldChar")
            separate_char.set(qn("w:fldCharType"), "separate")
            separate.append(separate_char)
            value_run = paragraph.add_run("1")
            value_run.font.size = Pt(9)
            end = paragraph.add_run()._r
            end_char = OxmlElement("w:fldChar")
            end_char.set(qn("w:fldCharType"), "end")
            end.append(end_char)
            changed_parts.add(part_name)
    return changed_parts


def _apply_paragraph_template(paragraph, sample_xml):
    target = paragraph._p
    current_ppr = target.find(qn("w:pPr"))
    if current_ppr is not None:
        target.remove(current_ppr)
    sample_ppr = sample_xml.find(qn("w:pPr"))
    if sample_ppr is not None:
        target.insert(0, deepcopy(sample_ppr))

    sample_run = sample_xml.find(qn("w:r"))
    if sample_run is None or not paragraph.runs:
        return
    sample_rpr = sample_run.find(qn("w:rPr"))
    target_run = paragraph.runs[0]._r
    current_rpr = target_run.find(qn("w:rPr"))
    if current_rpr is not None:
        target_run.remove(current_rpr)
    if sample_rpr is not None:
        target_run.insert(0, deepcopy(sample_rpr))


def _set_semantic_heading(paragraph, level):
    level = max(1, min(9, int(level or 1)))
    try:
        paragraph.style = f"Heading {level}"
    except KeyError:
        pass
    ppr = paragraph._p.get_or_add_pPr()
    outline = ppr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        ppr.append(outline)
    outline.set(qn("w:val"), str(level - 1))


def _paragraph_has_numbering(paragraph):
    if paragraph._p.xpath("./w:pPr/w:numPr"):
        return True
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        ppr = style.element.pPr
        if ppr is not None and ppr.find(qn("w:numPr")) is not None:
            return True
        style = style.base_style
    return False


def _heading_text_for_paragraph(title, paragraph, level):
    text = str(title or "").strip()
    if not _paragraph_has_numbering(paragraph):
        return text
    if int(level or 1) == 1:
        return re.sub(r"^第[一二三四五六七八九十百0-9]+章\s*", "", text)
    return re.sub(r"^\d+(?:\.\d+){1,8}\s*", "", text)


class DocxBuilder:
    def __init__(self, workspace_root, job_store, content_store):
        self.workspace_root = Path(workspace_root).resolve()
        self.job_store = job_store
        self.content_store = content_store

    def generate(self, job_id, output_name=""):
        job = self.job_store.get(job_id)
        missing = self.content_store.get_missing(job_id)
        if not missing["complete"]:
            raise ValueError(f"Required document sections are missing: {missing['missing']}")
        if not job.get("outline"):
            raise ValueError("Confirm the document outline before generation")
        sections = self.content_store.all_sections(job_id)
        contract = self._validate_generation_gate(job, sections)
        self.job_store.update(job_id, status="generating")

        version = self.job_store.next_version(job_id)
        revision_dir = self.job_store.job_root(job_id) / "revisions" / f"v{version:03d}"
        revision_dir.mkdir(parents=True, exist_ok=False)
        try:
            return self._generate_into(job, job_id, version, revision_dir, sections, contract, output_name)
        except BaseException:
            # Never leave a half-built revision directory behind; orphan dirs block later versions.
            shutil.rmtree(revision_dir, ignore_errors=True)
            self.job_store.update(job_id, status="drafting")
            raise

    def _generate_into(self, job, job_id, version, revision_dir, sections, contract, output_name):
        internal_docx = revision_dir / "document.docx"
        template_path = Path(job["template"]["path"])
        shutil.copyfile(template_path, internal_docx)

        document = Document(internal_docx)
        requirements = dict(job.get("requirements") or {})
        self._replace_cover(
            document,
            requirements.get("title") or requirements.get("topic") or job.get("title"),
            requirements,
        )
        table_samples = [
            (len(table.columns), len(table.rows), deepcopy(table._tbl))
            for table in document.tables
        ]
        self._rewrite_tables(document, sections, contract)
        self._rewrite_sections(document, sections, table_samples)
        _continue_page_numbering(document)
        changed_footer_parts = _normalize_floating_footer_page_fields(document)
        _request_field_update(document)
        document.save(internal_docx)
        allowed_changed_parts = {
            "word/document.xml",
            "word/settings.xml",
            *changed_footer_parts,
        }
        preserve_package_parts(
            template_path,
            internal_docx,
            allowed_changed_parts,
        )

        docx_sha256 = sha256_file(internal_docx)
        current = self._current_revision_or_none(job)
        if current is not None and current.get("docx_sha256") == docx_sha256:
            raise ValueError(
                "NO_CONTENT_CHANGE: the generated DOCX is byte-identical to "
                f"v{int(current.get('version') or 0):03d}, so a new version was not created. "
                "Regenerating cannot fix QA findings by itself. Change section content "
                "(document_content.upsert_section), the contract, or use docx_edit on the current "
                "revision; if a QA finding is environmental (fonts, Word, vision model), report it "
                "to the user instead."
            )

        output_dir = self.workspace_root / "outputs" / job["slug"]
        output_dir.mkdir(parents=True, exist_ok=True)
        filename = str(output_name or f"{job['slug']}-v{version:03d}.docx")
        if not filename.lower().endswith(".docx"):
            filename += ".docx"
        filename = Path(filename).name
        published = output_dir / filename
        if published.exists():
            raise FileExistsError(
                f"Published document version already exists: {published}. "
                "Pass a different output_name to docx_generate instead of deleting user files."
            )
        # The manifest is the single source of truth for revision metadata;
        # job.json only keeps a slim index pointing at manifest_path.
        manifest_path = revision_dir / "artifact-manifest.json"
        manifest = {
            "version": version,
            "base": "template",
            "template_sha256": job["template"]["sha256"],
            "docx_path": str(internal_docx),
            "published_path": None,
            "pending_published_path": str(published),
            "docx_sha256": docx_sha256,
            "package_parts": package_part_hashes(internal_docx),
            "allowed_changed_parts": sorted(allowed_changed_parts),
            "created_at": _now_iso(),
            "qa_passed": False,
            "delivery_ready": False,
            "manifest_path": str(manifest_path),
        }
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        self.job_store.add_revision(job_id, manifest)
        self.job_store.update(job_id, status="verifying")
        return {
            "ok": True,
            "job_id": job_id,
            "version": version,
            "docx_path": str(internal_docx),
            "published_path": None,
            "pending_published_path": str(published),
            "delivery_ready": False,
            "template_unchanged": sha256_file(template_path) == job["template"]["sha256"],
            "artifacts": [str(manifest_path)],
            "next_action": "docx_verify",
            "instruction": "Not a deliverable yet: run docx_verify(mode='all') on this version.",
        }

    @staticmethod
    def _current_revision_or_none(job):
        current = job.get("current_revision")
        if current is None:
            return None
        for item in job.get("revisions", []):
            if int(item.get("version", -1)) == int(current):
                return item
        return None

    @staticmethod
    def _replace_cover(document, title, requirements):
        if not title:
            return
        paragraphs = list(document.paragraphs)
        first_heading = next(
            (index for index, paragraph in enumerate(paragraphs) if _styled_heading_level(paragraph)),
            None,
        )
        if first_heading is None:
            first_heading = next(
                (index for index, paragraph in enumerate(paragraphs) if _is_heading(paragraph)),
                len(paragraphs),
            )
        nonempty = [paragraph for paragraph in paragraphs[:first_heading] if paragraph.text.strip()]
        if nonempty:
            _replace_paragraph_text(nonempty[0], title)

        brand_decision = str(
            requirements.get("brand_decision") or requirements.get("brand_info") or ""
        ).strip().lower()
        company_name = str(requirements.get("company_name") or "").strip()
        for paragraph in paragraphs[:first_heading]:
            if "编制单位" not in paragraph.text:
                continue
            if company_name:
                _replace_paragraph_text(paragraph, f"编制单位：{company_name}")
            elif brand_decision in {"delete", "remove", "删除"}:
                _replace_paragraph_text(paragraph, "")

    def _rewrite_sections(self, document, sections, table_samples=None):
        paragraphs = list(document.paragraphs)
        all_headings = [paragraph for paragraph in paragraphs if _styled_heading_level(paragraph)]
        if not all_headings:
            all_headings = [paragraph for paragraph in paragraphs if _is_heading(paragraph)]
        first_heading_index = paragraphs.index(all_headings[0]) if all_headings else 0
        body_sample = next(
            (paragraph for paragraph in paragraphs[first_heading_index + 1 :] if _is_body_sample(paragraph)),
            None,
        )
        body_sample_xml = deepcopy(body_sample._p) if body_sample is not None else None
        heading_samples = {}
        for paragraph in all_headings:
            heading_samples.setdefault(_heading_level(paragraph), deepcopy(paragraph._p))

        # Clear the old paragraph bodies while leaving tables and opaque objects in place.
        for index, heading in enumerate(all_headings):
            next_heading = all_headings[index + 1] if index + 1 < len(all_headings) else None
            parent = heading._p.getparent()
            children = list(parent)
            start_index = children.index(heading._p)
            end_index = children.index(next_heading._p) if next_heading is not None else len(children)
            for child in children[start_index + 1 : end_index]:
                if child.tag == qn("w:p"):
                    parent.remove(child)

        matched = min(len(all_headings), len(sections))
        last_anchor = None
        for index in range(matched):
            heading = all_headings[index]
            section = sections[index]
            sample_xml = self._heading_sample(heading_samples, int(section.get("level") or 1))
            if sample_xml is not None:
                _apply_paragraph_template(heading, sample_xml)
            _set_semantic_heading(heading, int(section.get("level") or 1))
            _replace_paragraph_text(
                heading,
                _heading_text_for_paragraph(
                    section["title"], heading, int(section.get("level") or 1)
                ),
            )
            last_anchor = self._insert_section_body(
                document,
                heading._p,
                heading._parent,
                section,
                body_sample_xml,
                table_samples or [],
            )

        if len(all_headings) > len(sections):
            for heading in reversed(all_headings[len(sections) :]):
                parent = heading._p.getparent()
                if parent is not None:
                    parent.remove(heading._p)

        if len(sections) > len(all_headings):
            if not all_headings:
                for section in sections:
                    heading = document.add_heading(section["title"], level=int(section.get("level") or 1))
                    last_anchor = self._insert_section_body(
                        document,
                        heading._p,
                        heading._parent,
                        section,
                        body_sample_xml,
                        table_samples or [],
                    )
                return
            anchor = last_anchor if last_anchor is not None else all_headings[-1]._p
            parent = all_headings[-1]._parent
            for section in sections[len(all_headings) :]:
                sample_xml = self._heading_sample(heading_samples, int(section.get("level") or 1))
                if sample_xml is None:
                    sample_xml = deepcopy(all_headings[-1]._p)
                heading_xml = deepcopy(sample_xml)
                anchor.addnext(heading_xml)
                heading = Paragraph(heading_xml, parent)
                _set_semantic_heading(heading, int(section.get("level") or 1))
                _replace_paragraph_text(
                    heading,
                    _heading_text_for_paragraph(
                        section["title"], heading, int(section.get("level") or 1)
                    ),
                )
                anchor = self._insert_section_body(
                    document,
                    heading_xml,
                    parent,
                    section,
                    body_sample_xml,
                    table_samples or [],
                )

    @staticmethod
    def _heading_sample(samples, level):
        if not samples:
            return None
        if level in samples:
            return samples[level]
        nearest = min(samples, key=lambda candidate: (abs(candidate - level), candidate))
        return samples[nearest]

    @classmethod
    def _insert_section_body(cls, document, anchor, parent, section, body_sample_xml, table_samples):
        blocks = _content_blocks(section.get("content"))
        if str(section.get("fact_status") or "").lower() == "assumption":
            visible_text = "\n".join(
                str(value) for block_type, value in blocks if block_type == "paragraph"
            )
            if not re.search(r"(?:【假设】|\[假设\]|假设说明|暂按|测算假设)", visible_text):
                blocks.insert(0, ("paragraph", "【假设】本节数据为用户确认的测算假设，需以最终资料为准。"))
        for block_type, value in blocks:
            if block_type == "table":
                anchor = cls._insert_content_table(
                    document,
                    anchor,
                    parent,
                    value,
                    table_samples,
                )
                continue
            if body_sample_xml is None:
                paragraph_xml = OxmlElement("w:p")
            else:
                paragraph_xml = deepcopy(body_sample_xml)
            anchor.addnext(paragraph_xml)
            paragraph = Paragraph(paragraph_xml, parent)
            _replace_paragraph_text(paragraph, value)
            anchor = paragraph_xml
        return anchor

    @staticmethod
    def _insert_content_table(document, anchor, parent, table_data, table_samples):
        headers = list(table_data.get("headers") or [])
        rows = [list(row) for row in table_data.get("rows") or []]
        values = ([headers] if headers else []) + rows
        if not values:
            return anchor
        columns = max(len(row) for row in values)
        matching = [sample for sample in table_samples if sample[0] == columns]
        if matching:
            _column_count, _row_count, sample_xml = min(matching, key=lambda item: item[1])
            table_xml = deepcopy(sample_xml)
            anchor.addnext(table_xml)
            table = Table(table_xml, parent)
        else:
            table = document.add_table(rows=1, cols=columns)
            try:
                table.style = "Table Grid"
            except KeyError:
                pass
            table_xml = table._tbl
            anchor.addnext(table_xml)
        while len(table.rows) < len(values):
            table._tbl.append(deepcopy(table.rows[-1]._tr))
        while len(table.rows) > len(values):
            table._tbl.remove(table.rows[-1]._tr)
        for row_index, row_values in enumerate(values):
            for column_index in range(columns):
                value = row_values[column_index] if column_index < len(row_values) else ""
                _table_set_text(table.rows[row_index].cells[column_index], value)
        return table_xml

    def _validate_generation_gate(self, job, sections):
        if not job.get("plan_confirmed"):
            raise ValueError("The user must confirm the document plan before generation")
        if job.get("pending_questions"):
            raise ValueError(
                "PENDING_QUESTIONS: resolve pending document requirements before generation. "
                "Save the user's answers with document_job.update_requirements("
                "requirements={...}, pending_questions=[])."
            )
        placeholder_titles = [
            section["id"]
            for section in sections
            if re.fullmatch(r"section-\d+", str(section.get("title") or "").strip())
        ]
        if placeholder_titles:
            raise ValueError(
                f"PLACEHOLDER_TITLES: outline sections still use placeholder titles: {placeholder_titles}. "
                "These ids would be rendered as literal headings. Set real Chinese titles in the outline "
                "(document_job.set_plan) or via document_content.upsert_section title before generation."
            )
        if job.get("source_candidates") and not job.get("source_confirmation_at"):
            raise ValueError("Workspace source candidates require explicit confirmation before generation")
        self._validate_literature_sources(job, sections)
        contract_path = job.get("template_contract_path")
        if not contract_path or not Path(contract_path).exists() or not job.get("contract_confirmed"):
            raise ValueError("A confirmed template contract is required before generation")
        contract = normalize_template_contract(
            json.loads(Path(contract_path).read_text(encoding="utf-8"))
        )
        requested = [table for section in sections for table in section.get("tables", [])]
        targeted = {str(item.get("target_element_id")): item for item in requested if item.get("target_element_id")}
        untargeted = [item for item in requested if not item.get("target_element_id")]
        table_items = {
            str(item.get("element_id")): item
            for item in list(contract.get("tables") or [])
            if isinstance(item, dict) and item.get("element_id")
        }
        contract_replacements = self._contract_table_replacements(contract)
        defaults = dict(contract.get("defaults") or {})
        spec = self._load_template_spec(job)
        for table in spec.get("tables", []):
            element_id = table["element_id"]
            action = str((table_items.get(element_id) or {}).get("action") or defaults.get("tables") or "confirm")
            if action == "confirm":
                raise ValueError(
                    f"Template table still requires user confirmation: {element_id}. "
                    "Resolve it with document_job.set_contract using "
                    f'{{"tables":[{{"element_id":"{element_id}","action":"preserve|delete|rewrite"}}]}}; '
                    "set_contract merges this decision with existing table decisions; then call "
                    "document_job.confirm_plan once. Do not use set_plan or set_outline for table decisions."
                )
            table_data = None
            if action in {"rewrite", "reuse_structure"}:
                table_data = targeted.get(element_id) or contract_replacements.get(element_id)
                if table_data is None:
                    if untargeted:
                        table_data = untargeted.pop(0)
                    else:
                        raise ValueError(
                            f"Replacement data is required for template table: {element_id}. "
                            "Attach it to the owning section via document_content.upsert_section "
                            f'tables=[{{"target_element_id":"{element_id}","headers":[...],"rows":[...]}}].'
                        )
                self._validate_table_columns(element_id, table, table_data)
            if action not in {"preserve", "rewrite", "reuse_structure", "delete"}:
                raise ValueError(f"Unsupported template contract action for {element_id}: {action}")
        return contract

    @staticmethod
    def _validate_table_columns(element_id, spec_table, table_data):
        expected = int(spec_table.get("columns") or 0)
        if not expected or not isinstance(table_data, dict):
            return
        headers = list(table_data.get("headers") or [])
        rows = [list(row) for row in table_data.get("rows") or []]
        all_rows = ([headers] if headers else []) + rows
        if not all_rows:
            return
        provided = max(len(row) for row in all_rows)
        if provided > expected:
            raise ValueError(
                f"TABLE_COLUMN_MISMATCH: template table {element_id} has {expected} columns but the "
                f"replacement data provides {provided}. Extra columns would break the template geometry "
                f"and fail verification. Provide at most {expected} columns per headers/rows entry, or "
                "set the table action to preserve."
            )

    def _validate_literature_sources(self, job, sections):
        if not is_literature_review_job(job):
            return
        job_id = job["id"]
        grounding = self.content_store.get_grounding_gaps(job_id)
        legal_ids = set(grounding["legal_source_ids"])
        if not legal_ids:
            raise ValueError(
                "SOURCE_GROUNDING_REQUIRED: literature reviews require confirmed workspace "
                "sources or web sources recorded with document_source.record_web before generation"
            )
        if grounding["gaps"]:
            invalid_sections = [item["section_id"] for item in grounding["gaps"]]
            raise ValueError(
                "SOURCE_GROUNDING_REQUIRED: literature-review sections lack grounded "
                f"source_ids: {invalid_sections}. Call document_content.get_grounding_gaps, "
                "then document_content.set_provenance for existing text; do not call set_outline "
                "or overwrite section content."
            )

    @staticmethod
    def _load_template_spec(job):
        path = job.get("template_spec_path")
        if not path or not Path(path).exists():
            raise ValueError("Analyze the DOCX template before generation")
        return json.loads(Path(path).read_text(encoding="utf-8"))

    @staticmethod
    def _contract_table_actions(contract, table_count):
        items = {
            str(item.get("element_id")): str(item.get("action") or "confirm")
            for item in list(contract.get("tables") or [])
            if isinstance(item, dict) and item.get("element_id")
        }
        default = str((contract.get("defaults") or {}).get("tables") or "confirm")
        return [items.get(f"body.tbl{index:04d}", default) for index in range(table_count)]

    @staticmethod
    def _contract_table_replacements(contract):
        replacements = {}
        for item in list(contract.get("tables") or []):
            if not isinstance(item, dict) or not item.get("element_id"):
                continue
            replacement = item.get("replacement_data")
            if isinstance(replacement, dict) and (
                replacement.get("headers") or replacement.get("rows")
            ):
                replacements[str(item["element_id"])] = dict(replacement)
        return replacements

    @classmethod
    def _rewrite_tables(cls, document, sections, contract):
        requested = [dict(table) for section in sections for table in section.get("tables", [])]
        targeted = {
            str(item.get("target_element_id")): item
            for item in requested
            if item.get("target_element_id")
        }
        untargeted = [item for item in requested if not item.get("target_element_id")]
        contract_replacements = cls._contract_table_replacements(contract)
        actions = cls._contract_table_actions(contract, len(document.tables))
        data_by_index = {}
        for index, action in enumerate(actions):
            if action not in {"rewrite", "reuse_structure"}:
                continue
            element_id = f"body.tbl{index:04d}"
            table_data = targeted.get(element_id)
            if table_data is None and untargeted:
                table_data = untargeted.pop(0)
            if table_data is None:
                table_data = contract_replacements.get(element_id)
            data_by_index[index] = table_data
        for index in range(len(document.tables) - 1, -1, -1):
            table = document.tables[index]
            element_id = f"body.tbl{index:04d}"
            action = actions[index]
            if action == "preserve":
                continue
            if action == "delete":
                table._tbl.getparent().remove(table._tbl)
                continue
            table_data = data_by_index.get(index)
            if table_data is None:
                raise ValueError(f"Replacement data is required for template table: {element_id}")
            headers = list(table_data.get("headers") or [])
            rows = [list(row) for row in table_data.get("rows") or []]
            all_rows = ([headers] if headers else []) + rows
            if not all_rows:
                continue
            column_count = max(len(row) for row in all_rows)
            while len(table.columns) < column_count:
                table.add_column(table.columns[-1].width)
            while len(table.rows) < len(all_rows):
                source_row = table.rows[-1]._tr
                table._tbl.append(deepcopy(source_row))
            while len(table.rows) > len(all_rows):
                row = table.rows[-1]._tr
                row.getparent().remove(row)
            for row_index, row_values in enumerate(all_rows):
                for column_index, value in enumerate(row_values):
                    if column_index < len(table.rows[row_index].cells):
                        _table_set_text(table.rows[row_index].cells[column_index], value)
