import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from lxml import etree

from yc_agents.documents.headings import structural_heading_level
from yc_agents.documents.ooxml import package_part_hashes, validate_docx_package


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
W = {"w": W_NS, "wp": WP_NS}


def _length(value):
    if value is None:
        return None
    try:
        return {"emu": int(value), "pt": round(float(value.pt), 4), "cm": round(float(value.cm), 4)}
    except (AttributeError, TypeError, ValueError):
        return value


def _enum(value):
    if value is None:
        return None
    return getattr(value, "name", str(value))


def _rgb(value):
    if value is None:
        return None
    try:
        return str(value.rgb) if value.rgb is not None else None
    except (AttributeError, ValueError):
        return None


def _style_chain(style):
    chain = []
    seen = set()
    current = style
    while current is not None and current.style_id not in seen:
        seen.add(current.style_id)
        chain.append(current)
        current = current.base_style
    return chain


def _first_value(objects, attribute):
    for obj in objects:
        if obj is None:
            continue
        value = getattr(obj, attribute, None)
        if value is not None:
            return value
    return None


def _rpr_elements(run, paragraph):
    elements = []
    direct = run._r.rPr
    if direct is not None:
        elements.append(direct)
    try:
        run_style = run.style
    except KeyError:
        run_style = None
    for style in _style_chain(run_style):
        rpr = style.element.rPr
        if rpr is not None:
            elements.append(rpr)
    for style in _style_chain(paragraph.style):
        rpr = style.element.rPr
        if rpr is not None:
            elements.append(rpr)
    return elements


def _font_names(elements):
    values = {}
    for key in ["ascii", "hAnsi", "eastAsia", "cs", "asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"]:
        attr = qn(f"w:{key}")
        for element in elements:
            fonts = element.find(qn("w:rFonts"))
            if fonts is not None and fonts.get(attr):
                values[key] = fonts.get(attr)
                break
    return values


def _spacing_raw(paragraph):
    elements = []
    if paragraph._p.pPr is not None:
        elements.append(paragraph._p.pPr)
    for style in _style_chain(paragraph.style):
        if style.element.pPr is not None:
            elements.append(style.element.pPr)
    result = {}
    for element in elements:
        spacing = element.find(qn("w:spacing"))
        if spacing is None:
            continue
        for key in ["before", "after", "line", "lineRule", "beforeAutospacing", "afterAutospacing"]:
            attr = spacing.get(qn(f"w:{key}"))
            if attr is not None and key not in result:
                result[key] = attr
    return result


def _indent_raw(paragraph):
    elements = []
    if paragraph._p.pPr is not None:
        elements.append(paragraph._p.pPr)
    for style in _style_chain(paragraph.style):
        if style.element.pPr is not None:
            elements.append(style.element.pPr)
    result = {}
    for element in elements:
        indent = element.find(qn("w:ind"))
        if indent is None:
            continue
        for key in ["left", "right", "firstLine", "hanging", "leftChars", "rightChars", "firstLineChars", "hangingChars"]:
            attr = indent.get(qn(f"w:{key}"))
            if attr is not None and key not in result:
                result[key] = attr
    return result


def _paragraph_format(paragraph):
    direct = paragraph.paragraph_format
    inherited = [style.paragraph_format for style in _style_chain(paragraph.style)]
    formats = [direct, *inherited]
    line_spacing = _first_value(formats, "line_spacing")
    if hasattr(line_spacing, "pt"):
        normalized_line_spacing = {"kind": "length", **_length(line_spacing)}
    elif line_spacing is not None:
        normalized_line_spacing = {"kind": "multiple", "value": float(line_spacing)}
    else:
        normalized_line_spacing = None
    return {
        "alignment": _enum(_first_value(formats, "alignment")),
        "left_indent": _length(_first_value(formats, "left_indent")),
        "right_indent": _length(_first_value(formats, "right_indent")),
        "first_line_indent": _length(_first_value(formats, "first_line_indent")),
        "space_before": _length(_first_value(formats, "space_before")),
        "space_after": _length(_first_value(formats, "space_after")),
        "line_spacing": normalized_line_spacing,
        "line_spacing_rule": _enum(_first_value(formats, "line_spacing_rule")),
        "keep_together": _first_value(formats, "keep_together"),
        "keep_with_next": _first_value(formats, "keep_with_next"),
        "page_break_before": _first_value(formats, "page_break_before"),
        "widow_control": _first_value(formats, "widow_control"),
        "raw_spacing": _spacing_raw(paragraph),
        "raw_indent": _indent_raw(paragraph),
        "numbering": _paragraph_numbering(paragraph),
    }


def _paragraph_numbering(paragraph):
    elements = []
    if paragraph._p.pPr is not None:
        elements.append(paragraph._p.pPr)
    for style in _style_chain(paragraph.style):
        if style.element.pPr is not None:
            elements.append(style.element.pPr)
    for element in elements:
        num_pr = element.find(qn("w:numPr"))
        if num_pr is None:
            continue
        ilvl = num_pr.find(qn("w:ilvl"))
        num_id = num_pr.find(qn("w:numId"))
        return {
            "level": int(ilvl.get(qn("w:val"))) if ilvl is not None else None,
            "num_id": int(num_id.get(qn("w:val"))) if num_id is not None else None,
        }
    return None


def _word_attributes(element, names):
    if element is None:
        return None
    values = {name: element.get(qn(f"w:{name}")) for name in names}
    return {key: value for key, value in values.items() if value is not None} or None


def _table_cell_margins(parent):
    if parent is None:
        return None
    result = {}
    for side in ["top", "left", "start", "bottom", "right", "end"]:
        node = parent.find(qn(f"w:{side}"))
        if node is not None:
            result[side] = _word_attributes(node, ["w", "type"])
    return result or None


def _borders(parent):
    if parent is None:
        return None
    result = {}
    for edge in ["top", "left", "start", "bottom", "right", "end", "insideH", "insideV"]:
        node = parent.find(qn(f"w:{edge}"))
        if node is not None:
            result[edge] = _word_attributes(node, ["val", "sz", "space", "color", "themeColor"])
    return result or None


def _font_format(run, paragraph):
    direct = run.font
    fonts = [direct]
    try:
        fonts.extend(style.font for style in _style_chain(run.style))
    except KeyError:
        pass
    fonts.extend(style.font for style in _style_chain(paragraph.style))
    rpr_elements = _rpr_elements(run, paragraph)
    return {
        "name": _first_value(fonts, "name"),
        "names": _font_names(rpr_elements),
        "size": _length(_first_value(fonts, "size")),
        "bold": _first_value(fonts, "bold"),
        "italic": _first_value(fonts, "italic"),
        "underline": _enum(_first_value(fonts, "underline")),
        "strike": _first_value(fonts, "strike"),
        "subscript": _first_value(fonts, "subscript"),
        "superscript": _first_value(fonts, "superscript"),
        "all_caps": _first_value(fonts, "all_caps"),
        "small_caps": _first_value(fonts, "small_caps"),
        "color": next((_rgb(font.color) for font in fonts if _rgb(font.color)), None),
    }


def _role_for(paragraph, index, first_nonempty):
    text = paragraph.text.strip()
    if not text:
        return "blank"
    if index == first_nonempty:
        return "title"
    heading_level = structural_heading_level(paragraph)
    if heading_level is not None:
        return f"heading_{min(heading_level, 4)}"
    if re.match(r"^(表|图)\s*[0-9一二三四五六七八九十]+", text):
        return "caption"
    return "body"


def _format_signature(element):
    paragraph = element.get("paragraph_format") or {}
    runs = element.get("runs") or []
    font = runs[0].get("effective_font", {}) if runs else {}
    signature = {
        "role": element.get("role"),
        "style_id": element.get("style_id"),
        "font": {
            "name": font.get("name"),
            "names": font.get("names"),
            "size": font.get("size"),
            "bold": font.get("bold"),
            "color": font.get("color"),
        },
        "paragraph": {
            "alignment": paragraph.get("alignment"),
            "first_line_indent": paragraph.get("first_line_indent"),
            "space_before": paragraph.get("space_before"),
            "space_after": paragraph.get("space_after"),
            "line_spacing": paragraph.get("line_spacing"),
            "line_spacing_rule": paragraph.get("line_spacing_rule"),
            "raw_indent": paragraph.get("raw_indent"),
            "raw_spacing": paragraph.get("raw_spacing"),
        },
    }
    return json.dumps(signature, ensure_ascii=False, sort_keys=True)


class DocxTemplateAnalyzer:
    def __init__(self, job_store):
        self.job_store = job_store

    def analyze(self, job_id):
        job = self.job_store.get(job_id)
        template = Path(job["template"]["path"])
        package = validate_docx_package(template)
        if package["sha256"] != job["template"]["sha256"]:
            raise ValueError("Template snapshot hash no longer matches the attachment")
        cached = self._cached_summary(job, package["sha256"])
        if cached is not None:
            # Re-analysis of the same template is a no-op: keep status and the
            # already answered/cleared pending_questions untouched.
            return cached
        self.job_store.update(job_id, status="analyzing_template")

        document = Document(template)
        elements = []
        nonempty = [index for index, paragraph in enumerate(document.paragraphs) if paragraph.text.strip()]
        first_nonempty = nonempty[0] if nonempty else -1
        for index, paragraph in enumerate(document.paragraphs):
            elements.append(self._paragraph_element(paragraph, index, first_nonempty, "word/document.xml"))

        headers = []
        footers = []
        for section_index, section in enumerate(document.sections):
            for kind, container, output in [
                ("header", section.header, headers),
                ("footer", section.footer, footers),
            ]:
                for index, paragraph in enumerate(container.paragraphs):
                    item = self._paragraph_element(
                        paragraph,
                        index,
                        0,
                        f"{kind}:section:{section_index}",
                    )
                    item["element_id"] = f"{kind}.s{section_index}.p{index:04d}"
                    output.append(item)

        tables = [self._table_spec(table, index) for index, table in enumerate(document.tables)]
        styles = [self._style_spec(style) for style in document.styles]
        sections = [self._section_spec(section, index) for index, section in enumerate(document.sections)]
        package_features = self._package_features(template)
        clusters = self._clusters(elements)
        spec = {
            "schema_version": 1,
            "template": {
                "path": str(template),
                "sha256": package["sha256"],
                "bytes": template.stat().st_size,
                "package_parts": package_part_hashes(template),
            },
            "package_report": package,
            "sections": sections,
            "styles": styles,
            "elements": elements,
            "tables": tables,
            "numbering": self._numbering_spec(template),
            "images": self._image_specs(template),
            "headers": headers,
            "footers": footers,
            "features": package_features,
            "format_clusters": clusters,
        }
        output_path = self.job_store.job_root(job_id) / "template" / "template-spec.json"
        self.job_store._write_json(output_path, spec)
        pending = self._pending_questions(package_features)
        self.job_store.update(
            job_id,
            status="waiting_requirements",
            template_spec_path=str(output_path),
            pending_questions=pending,
        )
        return {
            "ok": True,
            "job_id": job_id,
            "template_spec_path": str(output_path),
            "template_sha256": package["sha256"],
            "paragraphs": len(elements),
            "tables": len(tables),
            "sections": len(sections),
            "headers": len(headers),
            "footers": len(footers),
            "format_clusters": clusters[:20],
            "unsupported_features": package_features["unsupported"],
            "pending_questions": pending,
        }

    QUERY_CHAR_BUDGET = 60000
    QUERY_TABLE_ROW_LIMIT = 30

    def _cached_summary(self, job, sha256):
        spec_path = job.get("template_spec_path")
        if not spec_path or not Path(spec_path).exists():
            return None
        try:
            spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return None
        if ((spec.get("template") or {}).get("sha256")) != sha256:
            return None
        return {
            "ok": True,
            "cached": True,
            "job_id": job["id"],
            "template_spec_path": str(spec_path),
            "template_sha256": sha256,
            "paragraphs": len(spec.get("elements", [])),
            "tables": len(spec.get("tables", [])),
            "sections": len(spec.get("sections", [])),
            "headers": len(spec.get("headers", [])),
            "footers": len(spec.get("footers", [])),
            "format_clusters": list(spec.get("format_clusters", []))[:20],
            "unsupported_features": (spec.get("features") or {}).get("unsupported", []),
            "pending_questions": list(job.get("pending_questions") or []),
            "instruction": (
                "Template already analyzed; reuse this cached spec via docx_template_query "
                "instead of re-analyzing."
            ),
        }

    def query(self, job_id, role="", element_id="", part="", limit=20, detail=False):
        job = self.job_store.get(job_id)
        spec_path = job.get("template_spec_path")
        if not spec_path or not Path(spec_path).exists():
            raise FileNotFoundError("Template has not been analyzed")
        spec = json.loads(Path(spec_path).read_text(encoding="utf-8"))
        requested_role = str(role or "").strip()
        requested_part = str(part or "").strip()
        role = requested_role.casefold()
        part = requested_part.replace("\\", "/").casefold()
        heading_group = role in {"heading", "headings"}
        part_aliases = {
            "body": "word/document.xml",
            "document": "word/document.xml",
            "document.xml": "word/document.xml",
        }
        part = part_aliases.get(part, part)
        if not any([role, element_id, part]):
            return {
                "sections": spec["sections"],
                "features": spec["features"],
                "format_clusters": spec["format_clusters"][:limit],
                "numbering": spec.get("numbering", {}),
                "images": spec.get("images", [])[:limit],
            }
        candidates = [
            *spec["elements"],
            *spec["headers"],
            *spec["footers"],
            *spec.get("tables", []),
            *spec.get("images", []),
        ]
        if heading_group:
            candidates = [
                item
                for item in candidates
                if str(item.get("role") or "").startswith("heading_")
            ]
        elif role:
            candidates = [item for item in candidates if item.get("role") == role]
        if element_id:
            candidates = [item for item in candidates if item.get("element_id") == element_id]
        if part:
            candidates = [
                item
                for item in candidates
                if part in str(item.get("part") or "").casefold()
            ]
        full_detail = bool(detail) or bool(element_id)
        selected = candidates[: max(1, min(int(limit), 100))]
        matches = []
        used_chars = 0
        truncated = False
        for item in selected:
            rendered = self._render_match(item, full_detail)
            size = len(json.dumps(rendered, ensure_ascii=False))
            if matches and used_chars + size > self.QUERY_CHAR_BUDGET:
                truncated = True
                break
            used_chars += size
            matches.append(rendered)
        result = {"matches": matches, "count": len(candidates), "detail": full_detail}
        normalized_query = {}
        if heading_group:
            normalized_query["role"] = "heading_*"
        elif requested_role and role != requested_role:
            normalized_query["role"] = role
        if requested_part and part != requested_part.replace("\\", "/").casefold():
            normalized_query["part"] = part
        if normalized_query:
            result["normalized_query"] = normalized_query
        if truncated or len(matches) < len(selected):
            result["truncated"] = True
            result["instruction"] = (
                "Output capped. Narrow the query (element_id / smaller limit) or fetch one "
                "element at a time with detail=true for full formatting."
            )
        elif not matches:
            available_roles = sorted(
                {
                    str(item.get("role"))
                    for item in [
                        *spec["elements"],
                        *spec["headers"],
                        *spec["footers"],
                        *spec.get("tables", []),
                        *spec.get("images", []),
                    ]
                    if item.get("role")
                }
            )
            result["instruction"] = (
                "No matches. Use role='heading' for every heading level, role='table' "
                "for tables, part='body' for word/document.xml, or fetch a known "
                "element_id. Available roles: "
                f"{', '.join(available_roles[:20]) or '(none)'}."
            )
        elif not full_detail:
            result["instruction"] = (
                "Compact view (text + role only). Fetch a single element_id or pass "
                "detail=true when exact formatting is needed."
            )
        return result

    def _render_match(self, item, full_detail):
        if not full_detail:
            compact = {
                key: item.get(key)
                for key in ("element_id", "element_type", "role", "part", "style_id", "style_name")
                if item.get(key) is not None
            }
            if item.get("text") is not None:
                text = str(item.get("text") or "")
                compact["text"] = text[:200]
                if len(text) > 200:
                    compact["text_truncated"] = True
            for key in ("rows", "columns", "style", "grid_widths_dxa", "bytes", "pixels", "format"):
                if item.get(key) is not None:
                    compact[key] = item.get(key)
            return compact
        if item.get("element_type") == "table":
            rendered = dict(item)
            rows = list(rendered.get("row_details") or [])
            if len(rows) > self.QUERY_TABLE_ROW_LIMIT:
                rendered["row_details"] = rows[: self.QUERY_TABLE_ROW_LIMIT]
                rendered["row_details_truncated"] = (
                    f"showing {self.QUERY_TABLE_ROW_LIMIT} of {len(rows)} rows"
                )
            return rendered
        return dict(item)

    def _paragraph_element(self, paragraph, index, first_nonempty, part):
        runs = [
            {
                "text": run.text[:500],
                "style_id": getattr(run.style, "style_id", None),
                "effective_font": _font_format(run, paragraph),
            }
            for run in paragraph.runs
        ]
        return {
            "element_id": f"body.p{index:04d}",
            "element_type": "paragraph",
            "part": part,
            "order": index,
            "text": paragraph.text[:2000],
            "role": _role_for(paragraph, index, first_nonempty),
            "style_id": paragraph.style.style_id,
            "style_name": paragraph.style.name,
            "style_chain": [style.style_id for style in _style_chain(paragraph.style)],
            "paragraph_format": _paragraph_format(paragraph),
            "runs": runs,
        }

    @staticmethod
    def _style_spec(style):
        base_style = getattr(style, "base_style", None)
        item = {
            "style_id": style.style_id,
            "name": style.name,
            "type": _enum(style.type),
            "based_on": base_style.style_id if base_style is not None else None,
            "builtin": bool(getattr(style, "builtin", False)),
        }
        if style.type in {WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER}:
            item["font"] = {
                "name": style.font.name,
                "size": _length(style.font.size),
                "bold": style.font.bold,
                "italic": style.font.italic,
                "color": _rgb(style.font.color),
            }
            rpr = style.element.rPr
            item["font"]["names"] = _font_names([rpr] if rpr is not None else [])
        return item

    @staticmethod
    def _section_spec(section, index):
        return {
            "index": index,
            "orientation": _enum(section.orientation),
            "page_width": _length(section.page_width),
            "page_height": _length(section.page_height),
            "top_margin": _length(section.top_margin),
            "bottom_margin": _length(section.bottom_margin),
            "left_margin": _length(section.left_margin),
            "right_margin": _length(section.right_margin),
            "gutter": _length(section.gutter),
            "header_distance": _length(section.header_distance),
            "footer_distance": _length(section.footer_distance),
            "start_type": _enum(section.start_type),
            "different_first_page_header_footer": section.different_first_page_header_footer,
        }

    @staticmethod
    def _table_spec(table, index):
        grid = table._tbl.tblGrid
        widths = []
        if grid is not None:
            widths = [column.get(qn("w:w")) for column in grid.findall(qn("w:gridCol"))]
        rows = []
        for row_index, row in enumerate(table.rows):
            cells = []
            for cell_index, cell in enumerate(row.cells):
                tc_pr = cell._tc.tcPr
                width = tc_pr.tcW.get(qn("w:w")) if tc_pr is not None and tc_pr.tcW is not None else None
                cells.append(
                    {
                        "column": cell_index,
                        "text": cell.text[:1000],
                        "width_dxa": width,
                        "vertical_alignment": _enum(cell.vertical_alignment),
                        "grid_span": int(tc_pr.gridSpan.val) if tc_pr is not None and tc_pr.gridSpan is not None else 1,
                        "vertical_merge": (
                            str(tc_pr.vMerge.val) if tc_pr is not None and tc_pr.vMerge is not None else None
                        ),
                        "margins": _table_cell_margins(
                            tc_pr.find(qn("w:tcMar")) if tc_pr is not None else None
                        ),
                        "shading": _word_attributes(
                            tc_pr.find(qn("w:shd")) if tc_pr is not None else None,
                            ["val", "fill", "color", "themeFill"],
                        ),
                        "borders": _borders(
                            tc_pr.find(qn("w:tcBorders")) if tc_pr is not None else None
                        ),
                        "no_wrap": tc_pr.find(qn("w:noWrap")) is not None if tc_pr is not None else False,
                    }
                )
            tr_pr = row._tr.trPr
            height = None
            height_rule = None
            if tr_pr is not None:
                tr_height = tr_pr.find(qn("w:trHeight"))
                if tr_height is not None:
                    height = tr_height.get(qn("w:val"))
                    height_rule = tr_height.get(qn("w:hRule"))
            rows.append(
                {
                    "index": row_index,
                    "height": height,
                    "height_rule": height_rule,
                    "repeat_header": tr_pr.find(qn("w:tblHeader")) is not None if tr_pr is not None else False,
                    "cant_split": tr_pr.find(qn("w:cantSplit")) is not None if tr_pr is not None else False,
                    "cells": cells,
                }
            )
        tbl_pr = table._tbl.tblPr
        return {
            "element_id": f"body.tbl{index:04d}",
            "element_type": "table",
            "role": "table",
            "part": "word/document.xml",
            "style": table.style.style_id if table.style is not None else None,
            "rows": len(table.rows),
            "columns": len(table.columns),
            "grid_widths_dxa": widths,
            "width": _word_attributes(tbl_pr.find(qn("w:tblW")) if tbl_pr is not None else None, ["w", "type"]),
            "indent": _word_attributes(tbl_pr.find(qn("w:tblInd")) if tbl_pr is not None else None, ["w", "type"]),
            "alignment": _word_attributes(tbl_pr.find(qn("w:jc")) if tbl_pr is not None else None, ["val"]),
            "layout": _word_attributes(tbl_pr.find(qn("w:tblLayout")) if tbl_pr is not None else None, ["type"]),
            "cell_margins": _table_cell_margins(
                tbl_pr.find(qn("w:tblCellMar")) if tbl_pr is not None else None
            ),
            "borders": _borders(
                tbl_pr.find(qn("w:tblBorders")) if tbl_pr is not None else None
            ),
            "row_details": rows,
        }

    @staticmethod
    def _numbering_spec(path):
        with zipfile.ZipFile(path) as package:
            if "word/numbering.xml" not in package.namelist():
                return {"abstract": [], "instances": []}
            root = etree.fromstring(package.read("word/numbering.xml"))
        abstract = []
        for node in root.xpath("//w:abstractNum", namespaces=W):
            levels = []
            for level in node.xpath("./w:lvl", namespaces=W):
                def value(name):
                    child = level.find(qn(f"w:{name}"))
                    return child.get(qn("w:val")) if child is not None else None

                ppr = level.find(qn("w:pPr"))
                rpr = level.find(qn("w:rPr"))
                levels.append(
                    {
                        "level": int(level.get(qn("w:ilvl")) or 0),
                        "start": value("start"),
                        "format": value("numFmt"),
                        "text": value("lvlText"),
                        "alignment": value("lvlJc"),
                        "paragraph_style": value("pStyle"),
                        "suffix": value("suff"),
                        "indent": _word_attributes(ppr.find(qn("w:ind")) if ppr is not None else None, ["left", "hanging", "firstLine"]),
                        "fonts": _font_names([rpr] if rpr is not None else []),
                    }
                )
            abstract.append(
                {
                    "abstract_num_id": int(node.get(qn("w:abstractNumId"))),
                    "multi_level_type": next(iter(node.xpath("./w:multiLevelType/@w:val", namespaces=W)), None),
                    "levels": levels,
                }
            )
        instances = []
        for node in root.xpath("//w:num", namespaces=W):
            abstract_id = node.find(qn("w:abstractNumId"))
            instances.append(
                {
                    "num_id": int(node.get(qn("w:numId"))),
                    "abstract_num_id": int(abstract_id.get(qn("w:val"))) if abstract_id is not None else None,
                }
            )
        return {"abstract": abstract, "instances": instances}

    @staticmethod
    def _image_specs(path):
        images = []
        with zipfile.ZipFile(path) as package:
            for name in sorted(item for item in package.namelist() if item.startswith("word/media/") and not item.endswith("/")):
                data = package.read(name)
                item = {
                    "element_id": f"image.{len(images):04d}",
                    "element_type": "image",
                    "role": "image",
                    "part": name,
                    "bytes": len(data),
                }
                try:
                    from PIL import Image
                    from io import BytesIO

                    with Image.open(BytesIO(data)) as image:
                        item["pixels"] = {"width": image.width, "height": image.height}
                        item["format"] = image.format
                except Exception:
                    pass
                images.append(item)
        return images

    @staticmethod
    def _package_features(path):
        counts = Counter()
        unsupported = set()
        with zipfile.ZipFile(path) as package:
            for name in package.namelist():
                lower = name.lower()
                if "/charts/" in lower:
                    unsupported.add("chart")
                if "/embeddings/" in lower:
                    unsupported.add("embedded_object")
                if "/diagrams/" in lower:
                    unsupported.add("smartart")
                if not name.endswith(".xml"):
                    continue
                try:
                    root = etree.fromstring(package.read(name))
                except etree.XMLSyntaxError:
                    continue
                counts["fields"] += len(root.xpath("//w:fldSimple | //w:instrText", namespaces=W))
                counts["bookmarks"] += len(root.xpath("//w:bookmarkStart", namespaces=W))
                counts["content_controls"] += len(root.xpath("//w:sdt", namespaces=W))
                counts["drawings"] += len(root.xpath("//w:drawing", namespaces=W))
                counts["textboxes"] += len(root.xpath("//w:txbxContent", namespaces=W))
                if root.xpath("//w:txbxContent", namespaces=W):
                    unsupported.add("textbox")
        return {**dict(counts), "unsupported": sorted(unsupported)}

    @staticmethod
    def _clusters(elements):
        grouped = {}
        for item in elements:
            if item.get("role") == "blank":
                continue
            signature = _format_signature(item)
            group = grouped.setdefault(signature, {"count": 0, "examples": [], "format": json.loads(signature)})
            group["count"] += 1
            if len(group["examples"]) < 3:
                group["examples"].append(item.get("text", "")[:120])
        values = sorted(grouped.values(), key=lambda item: item["count"], reverse=True)
        for index, item in enumerate(values):
            item["cluster_id"] = f"fmt_{index:03d}"
        return values

    @staticmethod
    def _pending_questions(features):
        questions = [
            "新文档的主题、主要读者和必须包含的数据是什么？",
            "是否允许在已确认的工作区资料不足时联网搜索？",
            "模板中的公司名称、Logo和免责声明需要保留还是替换？",
        ]
        if features.get("unsupported"):
            questions.append(
                "模板包含暂不自动编辑的复杂对象，遇到旧业务内容时应保留、删除还是改用新图片？"
            )
        return questions
