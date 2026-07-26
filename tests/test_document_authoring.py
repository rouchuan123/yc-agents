import json
import os
import sys
import tempfile
import types
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from yc_agents.cli.commands import parse_cli_input
from yc_agents.documents.analyzer import DocxTemplateAnalyzer
from yc_agents.documents.attachments import AttachmentManager
from yc_agents.documents.broker import ExecutionBroker
from yc_agents.documents.builder import DocxBuilder, _heading_text_for_paragraph
from yc_agents.documents.content import DocumentContentStore
from yc_agents.documents.editor import DocxEditor
from yc_agents.documents.jobs import DocumentJobStore
from yc_agents.documents.ooxml import package_part_hashes, sha256_file
from yc_agents.documents.sources import DocumentSourceService
from yc_agents.documents.vision import VisionQAService
from yc_agents.documents.verifier import DocxVerifier
from yc_agents.documents.word_renderer import export_word_pdf
from yc_agents.tools.document_job import DocumentJobTool
from yc_agents.tools.document_content import DocumentContentTool
from yc_agents.tools.docx_verify import DocxVerifyTool


def make_template(path):
    document = Document()
    section = document.sections[0]
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    section.header.paragraphs[0].text = "示例公司"
    section.footer.paragraphs[0].text = "内部资料"

    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.add_run("旧项目可行性研究报告")

    heading = document.add_heading("第一章 项目概况", level=1)
    heading.paragraph_format.space_before = Pt(12)
    body = document.add_paragraph()
    body.style = document.styles["Normal"]
    body.paragraph_format.first_line_indent = Cm(0.74)
    body.paragraph_format.line_spacing = Pt(22)
    body.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    run = body.add_run("旧项目位于测试地区，需要替换为新内容。")
    run.font.name = "Times New Roman"
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(12)
    body._p.get_or_add_pPr().get_or_add_ind().set(qn("w:firstLineChars"), "200")

    document.add_heading("第二章 建设方案", level=1)
    document.add_paragraph("旧建设方案正文。")
    table = document.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    for index, text in enumerate(["设备", "数量", "负责人"]):
        table.rows[0].cells[index].text = text
    for index, text in enumerate(["旧设备", "1", "张三"]):
        table.rows[1].cells[index].text = text
    document.save(path)


@pytest.fixture
def document_workspace(tmp_path):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    session_path = workspace / ".ycore" / "sessions" / "session-test"
    session_path.mkdir(parents=True)
    template = tmp_path / "finished template.docx"
    make_template(template)
    attachments = AttachmentManager(session_path)
    attachment = attachments.import_file(template, role="template")
    jobs = DocumentJobStore(workspace, "session-test")
    job = jobs.create(attachment, title="新项目报告")
    return workspace, template, attachments, jobs, job


def test_cli_attachment_and_document_commands_parse_paths_with_spaces():
    assert parse_cli_input('/attach template "E:\\demo-docx\\a b.docx"').action == "attach"
    assert parse_cli_input("/attachments").action == "attachments"
    assert parse_cli_input("/detach att_1").content == "att_1"
    assert parse_cli_input("/document status").action == "document_status"
    assert parse_cli_input("/document rollback v2").content == "v2"


def test_attachment_manager_snapshots_template_and_rejects_docm(tmp_path):
    session = tmp_path / "session"
    template = tmp_path / "sample.docx"
    make_template(template)
    original_hash = sha256_file(template)
    manager = AttachmentManager(session)

    record = manager.import_file(template, role="template")

    assert record["sha256"] == original_hash
    assert Path(record["snapshot_path"]).read_bytes() == template.read_bytes()
    assert manager.get(record["id"])["role"] == "template"
    duplicate = manager.import_file(template, role="template")
    assert duplicate["id"] == record["id"]
    assert len(manager.list()) == 1
    macro = tmp_path / "unsafe.docm"
    macro.write_bytes(template.read_bytes())
    with pytest.raises(ValueError, match="docm"):
        manager.import_file(macro, role="template")


def test_attachment_manager_rejects_zip_path_traversal(tmp_path):
    unsafe = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(unsafe, "w") as package:
        package.writestr("[Content_Types].xml", "<Types />")
        package.writestr("word/document.xml", "<document />")
        package.writestr("../escape", "x")
    with pytest.raises(ValueError, match="Unsafe"):
        AttachmentManager(tmp_path / "session").import_file(unsafe, role="template")


def test_document_job_tool_returns_existing_attachments_and_auto_creates_job(tmp_path):
    workspace = tmp_path / "workspace"
    session_path = workspace / ".ycore" / "sessions" / "session-auto"
    session_path.mkdir(parents=True)
    template = tmp_path / "template.docx"
    make_template(template)
    attachments = AttachmentManager(session_path)
    attachment = attachments.import_file(template, role="template")
    jobs = DocumentJobStore(workspace, "session-auto")
    tool = DocumentJobTool(jobs, attachments)

    initial = tool.run("get_active")

    assert initial["job"] is None
    assert initial["attachments"] == [
        {
            "id": attachment["id"],
            "role": "template",
            "name": "template.docx",
            "suffix": ".docx",
            "bytes": attachment["bytes"],
            "sha256": attachment["sha256"],
            "created_at": attachment["created_at"],
        }
    ]

    created = tool.run("create", title="自动创建任务")

    assert created["selected_attachment"]["id"] == attachment["id"]
    assert created["job"]["title"] == "自动创建任务"
    assert tool.run("get_active")["job"]["id"] == created["job"]["id"]


def test_document_job_tool_uses_active_job_when_job_id_is_omitted(tmp_path):
    workspace = tmp_path / "workspace"
    session_path = workspace / ".ycore" / "sessions" / "session-active"
    session_path.mkdir(parents=True)
    template = tmp_path / "template.docx"
    make_template(template)
    attachments = AttachmentManager(session_path)
    attachment = attachments.import_file(template, role="template")
    jobs = DocumentJobStore(workspace, "session-active")
    job = jobs.create(attachment, title="活跃任务")
    tool = DocumentJobTool(jobs, attachments)

    result = tool.run(
        "set_contract",
        contract={
            "confirm": [
                {"element_id": "body.tbl0000", "decision": "rewrite"}
            ]
        },
    )

    assert result["job"]["id"] == job["id"]
    assert result["contract"]["tables"] == [
        {"element_id": "body.tbl0000", "action": "rewrite"}
    ]
    assert "confirm" not in result["contract"]


def test_document_job_high_frequency_operations_return_lite_view(document_workspace):
    _workspace, _template, attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    tool = DocumentJobTool(jobs, attachments)

    fetched = tool.run("get", job_id=job["id"])

    # The lite view keeps decision fields only: no outline, requirements or
    # per-revision dumps riding along on every call.
    lite = fetched["job"]
    assert lite["id"] == job["id"]
    assert lite["status"] == "waiting_requirements"
    assert lite["pending_question_count"] >= 3
    assert len(lite["pending_questions_head"]) == 3
    assert lite["plan_confirmed"] is False
    assert lite["current_revision"] is None
    assert lite["revision_count"] == 0
    assert lite["unresolved_confirm_count"] == 0  # no contract recorded yet
    for heavy in ("outline", "requirements", "revisions", "source_candidates"):
        assert heavy not in lite

    answered = tool.run(
        "update_requirements",
        job_id=job["id"],
        requirements={"topic": "新项目"},
        pending_questions=[],
    )
    undecided = tool.run("set_contract", job_id=job["id"], contract={"tables": []})
    # The spec table defaults to confirm, so the ledger counts one open item.
    assert undecided["job"]["unresolved_confirm_count"] == 1
    contracted = tool.run(
        "set_contract",
        job_id=job["id"],
        contract={"tables": [{"element_id": "body.tbl0000", "action": "preserve"}], "unresolved": []},
    )
    assert contracted["job"]["unresolved_confirm_count"] == 0
    planned = tool.run(
        "set_plan",
        job_id=job["id"],
        outline={"sections": [{"id": "s1", "title": "第一章 概况"}]},
    )
    confirmed = tool.run("confirm_plan", job_id=job["id"])

    # set_plan echoes the canonical outline so the model can proof-read it.
    assert planned["outline"]["sections"][0]["title"] == "第一章 概况"
    assert "outline" not in planned["job"]
    # get_outline fetches the full outline on demand instead of every response.
    outline = tool.run("get_outline", job_id=job["id"])
    assert outline["ok"] is True
    assert outline["plan_confirmed"] is True
    assert outline["outline"]["sections"][0]["id"] == "s1"

    for response in (fetched, answered, contracted, confirmed, tool.run("get_active")):
        assert len(json.dumps(response, ensure_ascii=False)) < 2048


def test_document_job_tool_reports_missing_active_job_clearly(tmp_path):
    workspace = tmp_path / "workspace"
    session_path = workspace / ".ycore" / "sessions" / "session-empty"
    session_path.mkdir(parents=True)
    tool = DocumentJobTool(
        DocumentJobStore(workspace, "session-empty"),
        AttachmentManager(session_path),
    )

    with pytest.raises(ValueError, match="No active document job"):
        tool.run("set_contract", contract={})


def test_template_analyzer_extracts_effective_chinese_formatting(document_workspace):
    _workspace, template, _attachments, jobs, job = document_workspace
    analyzer = DocxTemplateAnalyzer(jobs)

    result = analyzer.analyze(job["id"])
    body = analyzer.query(job["id"], role="body", limit=20, detail=True)["matches"]
    compact_body = analyzer.query(job["id"], role="body", limit=20)["matches"]
    table = analyzer.query(job["id"], element_id="body.tbl0000")["matches"][0]
    overview = analyzer.query(job["id"])
    matching = next(item for item in body if "旧项目位于" in item["text"])
    compact_matching = next(item for item in compact_body if "旧项目位于" in item["text"])

    # Role queries default to the compact projection to protect the context window.
    assert "paragraph_format" not in compact_matching
    assert "runs" not in compact_matching
    assert compact_matching["element_id"] == matching["element_id"]

    assert result["sections"] == 1
    assert result["tables"] == 1
    assert matching["paragraph_format"]["raw_indent"]["firstLineChars"] == "200"
    assert matching["paragraph_format"]["raw_spacing"]["line"] == "440"
    assert matching["runs"][0]["effective_font"]["names"]["eastAsia"] == "宋体"
    assert matching["runs"][0]["effective_font"]["size"]["pt"] == 12.0
    assert table["role"] == "table"
    assert table["columns"] == 3
    assert len(table["grid_widths_dxa"]) == 3
    assert "numbering" in overview
    assert sha256_file(template) == jobs.get(job["id"])["template"]["sha256"]


def test_content_generation_revision_and_rollback_are_immutable(document_workspace):
    workspace, template, _attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(
        job["id"],
        {
            "tables": [{"element_id": "body.tbl0000", "action": "reuse_structure"}],
            "unresolved": [],
        },
    )
    content = DocumentContentStore(jobs)
    content.set_outline(
        job["id"],
        {
            "sections": [
                {"id": "overview", "title": "第一章 新项目概况"},
                {"id": "plan", "title": "第二章 新建设方案"},
            ]
        },
    )
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "overview", "第一章 新项目概况", "新项目位于唐山市。\n\n本项目面向普通用户。")
    content.upsert_section(
        job["id"],
        "plan",
        "第二章 新建设方案",
        "采用新的建设方案。",
        tables=[{"headers": ["设备", "数量", "负责人"], "rows": [["充电桩", "10", "李四"]]}],
    )
    original_hash = sha256_file(template)
    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])

    assert generated["version"] == 1
    assert generated["published_path"] is None
    assert generated["delivery_ready"] is False
    assert not Path(generated["pending_published_path"]).exists()
    assert sha256_file(template) == original_hash
    generated_doc = Document(generated["docx_path"])
    assert any("新项目位于唐山市" in paragraph.text for paragraph in generated_doc.paragraphs)
    assert generated_doc.tables[0].rows[1].cells[0].text == "充电桩"
    template_parts = package_part_hashes(template)
    generated_parts = package_part_hashes(generated["docx_path"])
    for name, metadata in template_parts.items():
        if name not in {"word/document.xml", "word/settings.xml"}:
            assert generated_parts[name]["sha256"] == metadata["sha256"]

    edited = DocxEditor(workspace, jobs).edit(
        job["id"],
        1,
        [
            {"operation": "replace_text", "old_text": "唐山市", "new_text": "北京市"},
            {"operation": "delete_table_column", "target": "body.tbl0000", "column": "负责人"},
        ],
    )
    assert edited["version"] == 2
    assert edited["published_path"] is None
    assert edited["delivery_ready"] is False
    assert not Path(edited["pending_published_path"]).exists()
    assert Path(generated["docx_path"]).exists()
    edited_doc = Document(edited["docx_path"])
    assert any("北京市" in paragraph.text for paragraph in edited_doc.paragraphs)
    assert len(edited_doc.tables[0].columns) == 2
    edited_parts = package_part_hashes(edited["docx_path"])
    base_parts = package_part_hashes(generated["docx_path"])
    for name, metadata in base_parts.items():
        if name not in {"word/document.xml", "word/settings.xml"}:
            assert edited_parts[name]["sha256"] == metadata["sha256"]
    with pytest.raises(ValueError, match="Revision conflict"):
        DocxEditor(workspace, jobs).edit(job["id"], 1, [])
    rolled_back = jobs.rollback(job["id"], 1)
    assert rolled_back["current_revision"] == 1


def test_builder_maps_outline_to_top_level_headings_and_uses_body_format(tmp_path):
    template = tmp_path / "nested-headings.docx"
    document = Document()
    document.add_paragraph("旧标题", style="Title")
    document.add_heading("一、第一章", level=1)
    document.add_paragraph("第一章正文")
    document.add_heading("二、第二章", level=1)
    document.add_paragraph("第二章正文")
    document.add_heading("三、第三章", level=1)
    document.add_heading("3.1 第三章子标题", level=2)
    document.add_paragraph("第三章正文")
    document.save(template)

    workspace = tmp_path / "workspace"
    session_path = workspace / ".ycore" / "sessions" / "nested"
    session_path.mkdir(parents=True)
    attachment = AttachmentManager(session_path).import_file(template, role="template")
    jobs = DocumentJobStore(workspace, "nested")
    job = jobs.create(attachment, title="新标题")
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新标题"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)
    sections = [
        {"id": f"s{index}", "title": f"第{index}章 新标题"}
        for index in range(1, 4)
    ]
    content.set_outline(job["id"], {"sections": sections})
    jobs.confirm_plan(job["id"])
    for index, section in enumerate(sections, start=1):
        content.upsert_section(job["id"], section["id"], section["title"], f"第{index}章新正文")

    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])
    result = Document(generated["docx_path"])
    headings = [paragraph for paragraph in result.paragraphs if paragraph.style.name == "Heading 1"]
    body = next(paragraph for paragraph in result.paragraphs if paragraph.text == "第3章新正文")

    assert [paragraph.text for paragraph in headings] == [section["title"] for section in sections]
    assert body.style.name == "Normal"


def test_numbered_heading_does_not_duplicate_manual_chapter_prefix():
    document = Document()
    paragraph = document.add_heading("旧标题", level=1)
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_numId().val = 1

    assert _heading_text_for_paragraph("第一章 项目概述", paragraph, 1) == "项目概述"


def test_outline_accepts_chapters_alias_and_preserves_nested_heading_levels(tmp_path):
    template = tmp_path / "nested-template.docx"
    document = Document()
    document.add_paragraph("旧标题", style="Title")
    document.add_heading("第一章 旧内容", level=1)
    document.add_paragraph("一级正文")
    document.add_heading("1.1 旧子章节", level=2)
    document.add_paragraph("二级正文")
    document.add_heading("第二章 旧内容", level=1)
    document.add_paragraph("第二章正文")
    document.save(template)

    workspace = tmp_path / "workspace"
    session_path = workspace / ".ycore" / "sessions" / "nested-tree"
    session_path.mkdir(parents=True)
    attachment = AttachmentManager(session_path).import_file(template, role="template")
    jobs = DocumentJobStore(workspace, "nested-tree")
    job = jobs.create(attachment, title="新标题")
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新标题"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)

    outline = content.set_outline(
        job["id"],
        {
            "chapters": [
                {
                    "id": "ch1",
                    "title": "第一章 新内容",
                    "children": [{"id": "ch1-1", "title": "1.1 新子章节"}],
                },
                {"id": "ch2", "title": "第二章 新内容"},
            ]
        },
    )

    assert "chapters" not in outline
    assert outline["sections"][0]["level"] == 1
    assert outline["sections"][0]["children"][0]["level"] == 2
    assert outline["sections"][0]["children"][0]["parent_id"] == "ch1"
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "ch1", "第一章 新内容", "一级新正文")
    content.upsert_section(job["id"], "ch1-1", "1.1 新子章节", "二级新正文")
    content.upsert_section(job["id"], "ch2", "第二章 新内容", "末章新正文")

    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])
    result = Document(generated["docx_path"])
    headings = [paragraph for paragraph in result.paragraphs if paragraph.text in {
        "第一章 新内容", "1.1 新子章节", "第二章 新内容"
    }]

    assert [paragraph.text for paragraph in headings] == ["第一章 新内容", "1.1 新子章节", "第二章 新内容"]
    assert [paragraph.style.name for paragraph in headings] == ["Heading 1", "Heading 2", "Heading 1"]


def test_set_outline_feedback_requires_reconfirmation_and_stops_repeated_upserts(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)
    tool = DocumentContentTool(content)

    result = tool.run(
        "set_outline",
        job["id"],
        outline={"sections": [{"id": "s1", "title": "第一章"}]},
    )

    assert result["requires_plan_confirmation"] is True
    assert result["next_action"] == "document_job.confirm_plan"
    with pytest.raises(ValueError, match="PLAN_NOT_CONFIRMED.*confirm_plan"):
        tool.run("upsert_section", job["id"], section_id="s1", title="第一章", content="正文")


def test_document_content_tool_returns_compact_section_summary(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)
    content.set_outline(
        job["id"],
        {"sections": [{"id": "s1", "title": "第一章"}, {"id": "s2", "title": "第二章"}]},
    )
    jobs.confirm_plan(job["id"])
    tool = DocumentContentTool(content)

    body = "这是一段不会回显到工具历史中的长正文。"
    result = tool.run(
        "upsert_section",
        job["id"],
        section_id="s1",
        title="第一章",
        content=body,
        tables=[{"headers": ["项目"], "rows": [["值"]]}],
    )

    assert result["section"]["characters"] == len(body)
    assert result["section"]["tables"] == 1
    assert "content" not in result["section"]
    assert result["remaining"] == 1
    assert content.get_section(job["id"], "s1")["content"].startswith("这是一段")


def test_content_rejects_unsourced_project_metrics(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "第一章"}]})
    jobs.confirm_plan(job["id"])

    with pytest.raises(ValueError, match="UNSOURCED_PROJECT_METRIC"):
        content.upsert_section(job["id"], "s1", "第一章", "项目总投资5000万元，建筑面积15000平方米。")
    with pytest.raises(ValueError, match="confirmed source_ids"):
        content.upsert_section(
            job["id"],
            "s1",
            "第一章",
            "项目年营收8000万元。",
            fact_status="grounded",
            source_ids=["src_fake"],
        )

    accepted = content.upsert_section(
        job["id"],
        "s1",
        "第一章",
        "用户确认项目占地20亩。",
        fact_status="user_provided",
    )
    assert accepted["ok"] is True

    content.set_outline(
        job["id"],
        {
            "assumptions": ["总投资暂按5000万元测算"],
            "sections": [{"id": "s1", "title": "第一章"}],
        },
    )
    jobs.confirm_plan(job["id"])
    with pytest.raises(ValueError, match="UNCONFIRMED_PROJECT_ASSUMPTION"):
        content.upsert_section(
            job["id"],
            "s1",
            "第一章",
            "假设总投资5000万元、占地20亩。",
            fact_status="assumption",
        )
    assumed = content.upsert_section(
        job["id"],
        "s1",
        "第一章",
        "假设总投资5000万元。",
        fact_status="assumption",
    )
    assert assumed["ok"] is True


def test_generation_requires_confirmed_plan_and_table_replacements(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(
        job["id"],
        {"tables": [{"element_id": "body.tbl0000", "action": "reuse_structure"}], "unresolved": []},
    )
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "第一章"}]})

    with pytest.raises(ValueError, match="PLAN_NOT_CONFIRMED.*confirm_plan"):
        content.upsert_section(job["id"], "s1", "第一章", "正文")

    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "s1", "第一章", "正文")
    with pytest.raises(ValueError, match="Replacement data is required"):
        DocxBuilder(workspace, jobs, content).generate(job["id"])


def test_generation_accepts_legacy_contract_table_replacement_data(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(
        job["id"],
        {
            "tables": [
                {
                    "element_id": "body.tbl0000",
                    "action": "rewrite",
                    "replacement_data": {
                        "rows": [["阶段", "说明"], ["当前", "兼容旧任务"]]
                    },
                }
            ],
            "unresolved": [],
        },
    )
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "第一章"}]})
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "s1", "第一章", "正文")

    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])
    result = Document(generated["docx_path"])

    assert result.tables[0].cell(0, 0).text == "阶段"
    assert result.tables[0].cell(1, 1).text == "兼容旧任务"


def test_literature_review_sections_require_recorded_sources(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    jobs.update(job["id"], title="室内定位文献综述", slug="室内定位文献综述")
    jobs.update_requirements(job["id"], {"topic": "室内定位"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "研究进展"}]})
    jobs.confirm_plan(job["id"])

    with pytest.raises(ValueError, match="SOURCE_GROUNDING_REQUIRED"):
        content.upsert_section(job["id"], "s1", "研究进展", "Chen等（2024）提出了新方法。")

    source = DocumentSourceService(_workspace, jobs).record_web(
        job["id"],
        {"url": "https://example.com/paper", "title": "Verified paper"},
    )
    result = content.upsert_section(
        job["id"],
        "s1",
        "研究进展",
        "该研究提出了新的室内定位方法。",
        source_ids=[source["id"]],
        fact_status="grounded",
    )

    assert result["section"]["source_ids"] == [source["id"]]


def test_literature_provenance_repair_preserves_existing_content(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    jobs.update_requirements(job["id"], {"topic": "室内定位"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "研究进展"}]})
    jobs.confirm_plan(job["id"])
    original = "已有正文不能在修复来源时被覆盖。"
    content.upsert_section(job["id"], "s1", "研究进展", original)
    jobs.update(job["id"], title="室内定位系统综述", slug="室内定位系统综述")
    source = DocumentSourceService(workspace, jobs).record_web(
        job["id"],
        {"url": "https://example.com/paper", "title": "Verified paper"},
    )

    tool = DocumentContentTool(content)
    gaps = tool.run("get_grounding_gaps", job["id"])
    repaired = tool.run(
        "set_provenance",
        job["id"],
        section_id="s1",
        source_ids=[source["id"]],
        fact_status="grounded",
        reason="修复旧章节来源",
    )

    assert gaps["gaps"][0]["section_id"] == "s1"
    assert "missing_source_ids" in gaps["gaps"][0]["reasons"]
    assert repaired["section"]["characters"] == len(original)
    assert content.get_section(job["id"], "s1")["content"] == original
    assert content.get_grounding_gaps(job["id"])["complete"] is True


def test_empty_upsert_cannot_erase_existing_section(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "第一章"}]})
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "s1", "第一章", "已有正文")

    with pytest.raises(ValueError, match="EMPTY_SECTION_OVERWRITE.*set_provenance"):
        content.upsert_section(job["id"], "s1", "第一章", "")

    assert content.get_section(job["id"], "s1")["content"] == "已有正文"


def test_empty_required_leaf_section_remains_missing(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "第一章"}]})
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "s1", "第一章", "")

    assert content.get_missing(job["id"])["missing"] == ["s1"]


def test_literature_review_table_only_section_also_requires_sources(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    jobs.update(job["id"], title="室内定位系统综述", slug="室内定位系统综述")
    jobs.update_requirements(job["id"], {"topic": "室内定位"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "方法对比"}]})
    jobs.confirm_plan(job["id"])

    with pytest.raises(ValueError, match="SOURCE_GROUNDING_REQUIRED"):
        content.upsert_section(
            job["id"],
            "s1",
            "方法对比",
            "",
            tables=[{"headers": ["方法"], "rows": [["LLM-Loc"]]}],
        )

    source = DocumentSourceService(workspace, jobs).record_web(
        job["id"],
        {"url": "https://example.com/paper", "title": "Verified paper"},
    )
    content.upsert_section(
        job["id"],
        "s1",
        "方法对比",
        "",
        source_ids=[source["id"]],
        fact_status="grounded",
        tables=[{"headers": ["方法"], "rows": [["Verified method"]]}],
    )

    assert content.get_grounding_gaps(job["id"])["complete"] is True


def test_generation_reaudits_existing_literature_review_sections(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "室内定位"}, pending_questions=[])
    jobs.set_contract(
        job["id"],
        {"tables": [{"element_id": "body.tbl0000", "action": "preserve"}], "unresolved": []},
    )
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "研究进展"}]})
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "s1", "研究进展", "未标注来源的旧章节。")
    source = DocumentSourceService(workspace, jobs).record_web(
        job["id"],
        {"url": "https://example.com/paper", "title": "Verified paper"},
    )
    assert source["id"].startswith("web_")
    jobs.update(job["id"], title="室内定位文献综述", slug="室内定位文献综述")

    with pytest.raises(ValueError, match="sections lack grounded source_ids.*s1"):
        DocxBuilder(workspace, jobs, content).generate(job["id"])


def test_contract_legacy_id_is_normalized_and_deletes_table(document_workspace):
    workspace, _template, attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    tool = DocumentJobTool(jobs, attachments)

    response = tool.run(
        "set_contract",
        job_id=job["id"],
        contract={
            "tables": [{"id": "body.tbl0000", "action": "delete"}],
            "unresolved": [],
            "confirmed": True,
        },
    )

    assert response["contract"]["tables"] == [
        {"element_id": "body.tbl0000", "action": "delete"}
    ]
    assert response["contract"]["confirmed"] is False
    assert response["requires_plan_confirmation"] is True
    assert response["next_action"] == "document_job.confirm_plan"
    assert jobs.get(job["id"])["plan_confirmed"] is False

    content = DocumentContentStore(jobs)
    content.set_outline(
        job["id"],
        {
            "sections": [
                {"id": "s1", "title": "第一章"},
                {"id": "s2", "title": "第二章"},
            ]
        },
    )
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "s1", "第一章", "第一章正文")
    content.upsert_section(job["id"], "s2", "第二章", "第二章正文")

    # Emulate the already-persisted contract from the real failed job. The builder
    # must remain backward compatible even before a later confirm_plan rewrites it.
    contract_path = Path(jobs.get(job["id"])["template_contract_path"])
    contract_path.write_text(
        json.dumps(
            {
                "tables": [{"id": "body.tbl0000", "action": "delete"}],
                "unresolved": [],
                "confirmed": True,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])
    result = Document(generated["docx_path"])

    assert result.tables == []


def test_contract_rejects_conflicting_legacy_and_canonical_ids(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace

    with pytest.raises(ValueError, match="Conflicting element_id and id"):
        jobs.set_contract(
            job["id"],
            {
                "tables": [
                    {
                        "id": "body.tbl0000",
                        "element_id": "body.tbl0001",
                        "action": "delete",
                    }
                ]
            },
        )


def test_contract_legacy_confirm_decision_merges_and_canonical_item_wins(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace

    jobs.set_contract(
        job["id"],
        {"tables": [{"element_id": "body.tbl0001", "action": "delete"}]},
    )
    jobs.set_contract(
        job["id"],
        {"confirm": [{"element_id": "body.tbl0000", "decision": "rewrite"}]},
    )
    jobs.set_contract(
        job["id"],
        {
            "tables": [{"element_id": "body.tbl0000", "action": "preserve"}],
            "confirm": [{"element_id": "body.tbl0000", "decision": "rewrite"}],
        },
    )

    contract = jobs.get_contract(job["id"])
    assert contract["tables"] == [
        {"element_id": "body.tbl0001", "action": "delete"},
        {"element_id": "body.tbl0000", "action": "preserve"},
    ]
    assert "confirm" not in contract


def test_contract_rejects_confirmed_as_an_action_with_valid_choices(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace

    with pytest.raises(ValueError, match=r"confirmed.*Use one of:.*rewrite"):
        jobs.set_contract(
            job["id"],
            {"confirm": [{"element_id": "body.tbl0000", "decision": "confirmed"}]},
        )


def test_contract_partial_updates_merge_and_identical_repeat_is_idempotent(document_workspace):
    _workspace, _template, attachments, jobs, job = document_workspace
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "第一章"}]})
    tool = DocumentJobTool(jobs, attachments)

    first = tool.run(
        "set_contract",
        job_id=job["id"],
        contract={
            "tables": [
                {"element_id": "body.tbl0000", "action": "delete"},
                {"element_id": "body.tbl0001", "action": "delete"},
                {"element_id": "body.tbl0002", "action": "delete"},
            ]
        },
    )
    partial = tool.run(
        "set_contract",
        job_id=job["id"],
        contract={"tables": [{"element_id": "body.tbl0003", "action": "preserve"}]},
    )

    assert first["contract_changed"] is True
    assert [item["element_id"] for item in partial["contract"]["tables"]] == [
        "body.tbl0000",
        "body.tbl0001",
        "body.tbl0002",
        "body.tbl0003",
    ]
    confirmed = tool.run("confirm_plan", job_id=job["id"])
    assert confirmed["next_action"] == "docx_generate"

    repeated = tool.run(
        "set_contract",
        job_id=job["id"],
        contract={
            "tables": [
                {"element_id": "body.tbl0000", "action": "delete"},
                {"element_id": "body.tbl0001", "action": "delete"},
                {"element_id": "body.tbl0002", "action": "delete"},
                {"element_id": "body.tbl0003", "action": "preserve"},
            ]
        },
    )

    assert repeated["contract_changed"] is False
    assert repeated["requires_plan_confirmation"] is False
    assert repeated["next_action"] == "docx_generate"
    assert repeated["job"]["plan_confirmed"] is True
    assert repeated["job"]["contract_confirmed"] is True
    assert repeated["contract"]["confirmed"] is True

    with pytest.raises(ValueError, match="CONTRACT_LOCKED"):
        tool.run(
            "replace_contract",
            job_id=job["id"],
            contract={"tables": [{"element_id": "body.tbl0003", "action": "preserve"}]},
        )
    unlocked = tool.run("unlock_contract", job_id=job["id"])
    assert unlocked["job"]["contract_locked"] is False
    replaced = tool.run(
        "replace_contract",
        job_id=job["id"],
        contract={"tables": [{"element_id": "body.tbl0003", "action": "preserve"}]},
    )
    assert replaced["contract_changed"] is True
    assert replaced["requires_plan_confirmation"] is True
    assert [item["element_id"] for item in replaced["contract"]["tables"]] == ["body.tbl0003"]


def test_editor_scan_does_not_materialize_unreferenced_headers(tmp_path):
    template = tmp_path / "no-headers.docx"
    document = Document()
    document.add_heading("第一章 测试", level=1)
    document.add_paragraph("原正文")
    document.save(template)
    before = Document(template)
    assert not before.sections[0]._sectPr.findall(qn("w:headerReference"))
    assert not before.sections[0]._sectPr.findall(qn("w:footerReference"))

    paragraphs = DocxEditor._all_paragraphs(before)

    assert any(paragraph.text == "原正文" for paragraph in paragraphs)
    assert not before.sections[0]._sectPr.findall(qn("w:headerReference"))
    assert not before.sections[0]._sectPr.findall(qn("w:footerReference"))


def test_document_sources_require_confirmation_before_ingest(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    reference = workspace / "市场调研.md"
    reference.write_text("唐山市充电站市场需求增长，计划建设十台充电桩。", encoding="utf-8")
    sources = DocumentSourceService(workspace, jobs, chunk_size=20, chunk_overlap=2)
    discovered = sources.discover(job["id"], query="市场调研")
    source_id = next(item["id"] for item in discovered["candidates"] if item["path"] == "市场调研.md")

    assert jobs.get(job["id"])["confirmed_sources"] == []
    sources.confirm(job["id"], [source_id])
    ingested = sources.ingest(job["id"])
    result = sources.search(job["id"], "充电站")

    assert ingested["sources"][0]["path"] == "市场调研.md"
    assert result["results"]
    assert result["results"][0]["source"] == source_id


def test_document_source_discovery_excludes_template_copy_unless_explicitly_requested(document_workspace):
    workspace, template, _attachments, jobs, job = document_workspace
    template_copy = workspace / template.name
    template_copy.write_bytes(template.read_bytes())
    (workspace / "独立参考资料.md").write_text("可用业务资料", encoding="utf-8")
    sources = DocumentSourceService(workspace, jobs)

    default = sources.discover(job["id"])
    explicit = sources.discover(job["id"], include_template=True)

    assert template.name not in {item["name"] for item in default["candidates"]}
    assert template.name in {item["name"] for item in explicit["candidates"]}


def test_document_source_discovery_ignores_word_lock_files(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    (workspace / "~$文献综述.docx").write_bytes(b"word-lock")
    (workspace / "文献综述.md").write_text("可用资料", encoding="utf-8")

    result = DocumentSourceService(workspace, jobs).discover(job["id"], query="文献综述")

    names = {item["name"] for item in result["candidates"]}
    assert "~$文献综述.docx" not in names
    assert "文献综述.md" in names


def test_execution_broker_scrubs_secrets_and_restricts_paths(tmp_path, monkeypatch):
    root = tmp_path / "job"
    root.mkdir()
    source = root / "input.docx"
    source.write_bytes(b"x")
    output = root / "out.pdf"
    broker = ExecutionBroker([root], [root], timeout_seconds=10)
    monkeypatch.setenv("DEEPSEEK_API_KEY", "do-not-leak")

    def fake_command(_key, _input, output_path):
        return [
            sys.executable,
            "-c",
            "import os,pathlib,sys; pathlib.Path(sys.argv[1]).write_bytes(b'ok'); print(os.getenv('DEEPSEEK_API_KEY','missing'))",
            str(output_path),
        ]

    monkeypatch.setattr(broker, "_build_command", fake_command)
    result = broker.run("word_export_pdf", source, output)

    assert result["ok"] is True
    assert "missing" in result["stdout"]
    assert "DEEPSEEK_API_KEY" not in result["environment_keys"]
    with pytest.raises(PermissionError):
        broker.run("word_export_pdf", source, tmp_path / "outside.pdf")


def test_execution_broker_decodes_utf8_helper_output(tmp_path, monkeypatch):
    root = tmp_path / "job"
    root.mkdir()
    source = root / "input.docx"
    source.write_bytes(b"x")
    output = root / "out.pdf"
    broker = ExecutionBroker([root], [root], timeout_seconds=10)

    def fake_command(_key, _input, output_path):
        return [
            sys.executable,
            "-c",
            "import pathlib,sys; pathlib.Path(sys.argv[1]).write_bytes(b'ok'); print('渲染完成')",
            str(output_path),
        ]

    monkeypatch.setattr(broker, "_build_command", fake_command)
    result = broker.run("word_export_pdf", source, output)

    assert result["ok"] is True
    assert "渲染完成" in result["stdout"]


def _passthrough_render_command(_key, _input, output_path):
    return [
        sys.executable,
        "-c",
        "import pathlib,sys; pathlib.Path(sys.argv[1]).write_bytes(b'ok')",
        str(output_path),
    ]


def test_execution_broker_declares_scope_honestly_without_claiming_enforcement(tmp_path, monkeypatch):
    root = tmp_path / "job"
    root.mkdir()
    source = root / "input.docx"
    source.write_bytes(b"x")
    output = root / "out.pdf"
    broker = ExecutionBroker([root], [root], timeout_seconds=10)
    monkeypatch.setattr(broker, "_build_command", _passthrough_render_command)

    result = broker.run("word_export_pdf", source, output)

    assert result["ok"] is True
    # broker 不是 OS 级沙箱：作用域只是声明，必须带 enforced=False，
    # 且不再用会被误读为已强制执行的旧字段名。
    assert result["enforced"] is False
    assert result["declared_write_scope"] == [str(Path(root).resolve())]
    assert result["declared_network_policy"] == "not_enforced_local_backend"
    assert "write_scope" not in result
    assert "network_policy" not in result
    assert result["security_degraded"] == []
    assert result["out_of_scope_writes"] == []
    assert result["job_object"] in {"active", "unavailable"}
    # Word 是 out-of-process COM：Job Object 只能包住 python 子进程树，
    # 结果字段要如实说明这个边界。
    assert "WINWORD" in result["job_object_note"]


def test_execution_broker_lifts_security_degraded_from_renderer_stdout(tmp_path, monkeypatch):
    root = tmp_path / "job"
    root.mkdir()
    source = root / "input.docx"
    source.write_bytes(b"x")
    output = root / "out.pdf"
    broker = ExecutionBroker([root], [root], timeout_seconds=10)

    def fake_command(_key, _input, output_path):
        payload = json.dumps(
            {"ok": True, "security_degraded": ["automation_security"]},
            ensure_ascii=False,
        )
        return [
            sys.executable,
            "-c",
            "import pathlib,sys; pathlib.Path(sys.argv[1]).write_bytes(b'ok'); print(sys.argv[2])",
            str(output_path),
            payload,
        ]

    monkeypatch.setattr(broker, "_build_command", fake_command)
    result = broker.run("word_export_pdf", source, output)

    assert result["ok"] is True
    assert result["security_degraded"] == ["automation_security"]


def test_execution_broker_records_out_of_scope_writes_without_deleting(tmp_path, monkeypatch):
    root = tmp_path / "job"
    root.mkdir()
    source = root / "input.docx"
    source.write_bytes(b"x")
    output = root / "out.pdf"
    broker = ExecutionBroker([root], [root], timeout_seconds=10)

    def fake_command(_key, _input, output_path):
        return [
            sys.executable,
            "-c",
            (
                "import pathlib,sys; out=pathlib.Path(sys.argv[1]); "
                "out.write_bytes(b'ok'); "
                "(out.parent / 'unexpected.tmp').write_bytes(b'stray')"
            ),
            str(output_path),
        ]

    monkeypatch.setattr(broker, "_build_command", fake_command)
    result = broker.run("word_export_pdf", source, output)

    stray = root / "unexpected.tmp"
    assert result["ok"] is True
    assert result["out_of_scope_writes"] == [str(stray)]
    # 仅记录不删除：复核是审计动作，不是清理动作。
    assert stray.exists()


def test_execution_broker_degrades_gracefully_when_job_object_fails(tmp_path, monkeypatch):
    root = tmp_path / "job"
    root.mkdir()
    source = root / "input.docx"
    source.write_bytes(b"x")
    output = root / "out.pdf"
    broker = ExecutionBroker([root], [root], timeout_seconds=10)
    monkeypatch.setattr(broker, "_build_command", _passthrough_render_command)

    def broken_limiter(_pid):
        raise RuntimeError("win32job is unavailable")

    monkeypatch.setattr(
        "yc_agents.documents.broker._create_job_limiter", broken_limiter
    )
    result = broker.run("word_export_pdf", source, output)

    # Job Object 是尽力而为的加固：失败绝不影响渲染本身。
    assert result["ok"] is True
    assert result["job_object"] == "unavailable"


def test_word_renderer_uses_disposable_copy_and_does_not_save_revision(tmp_path, monkeypatch):
    source = tmp_path / "revision.docx"
    source.write_bytes(b"immutable-revision")
    output = tmp_path / "qa" / "document.pdf"
    opened_paths = []

    class FakeFields:
        def Update(self):
            return None

    class FakeDocument:
        Fields = FakeFields()
        TablesOfContents = []

        def ExportAsFixedFormat(self, **kwargs):
            Path(kwargs["OutputFileName"]).write_bytes(b"pdf")

        def Save(self):
            raise AssertionError("immutable revision must not be saved")

        def Close(self, SaveChanges=False):
            assert SaveChanges is False

    class FakeDocuments:
        def Open(self, path, **_kwargs):
            opened_paths.append(Path(path))
            assert Path(path) != source
            assert Path(path).read_bytes() == source.read_bytes()
            return FakeDocument()

    class FakeWord:
        Documents = FakeDocuments()
        Visible = True
        DisplayAlerts = 1
        AutomationSecurity = 0
        Options = types.SimpleNamespace(UpdateLinksAtOpen=True)

        def Quit(self):
            return None

    pythoncom = types.ModuleType("pythoncom")
    pythoncom.CoInitialize = lambda: None
    pythoncom.CoUninitialize = lambda: None
    win32com = types.ModuleType("win32com")
    client = types.ModuleType("win32com.client")
    client.DispatchEx = lambda _name: FakeWord()
    win32com.client = client
    monkeypatch.setitem(sys.modules, "pythoncom", pythoncom)
    monkeypatch.setitem(sys.modules, "win32com", win32com)
    monkeypatch.setitem(sys.modules, "win32com.client", client)

    result = export_word_pdf(source, output)

    assert result["ok"] is True
    assert result["security_degraded"] == []
    assert source.read_bytes() == b"immutable-revision"
    assert output.read_bytes() == b"pdf"
    assert opened_paths and not opened_paths[0].exists()


def test_word_renderer_reports_security_degradation_instead_of_swallowing(tmp_path, monkeypatch):
    source = tmp_path / "revision.docx"
    source.write_bytes(b"immutable-revision")
    output = tmp_path / "qa" / "document.pdf"

    class FakeFields:
        def Update(self):
            return None

    class FakeDocument:
        Fields = FakeFields()
        TablesOfContents = []

        def ExportAsFixedFormat(self, **kwargs):
            Path(kwargs["OutputFileName"]).write_bytes(b"pdf")

        def Close(self, SaveChanges=False):
            assert SaveChanges is False

    class FakeDocuments:
        def Open(self, path, **_kwargs):
            return FakeDocument()

    class FakeWord:
        Documents = FakeDocuments()

        def __setattr__(self, name, value):
            if name == "AutomationSecurity":
                raise OSError("AutomationSecurity is not supported")
            object.__setattr__(self, name, value)

        @property
        def Options(self):
            raise OSError("Options are unavailable")

        def Quit(self):
            return None

    pythoncom = types.ModuleType("pythoncom")
    pythoncom.CoInitialize = lambda: None
    pythoncom.CoUninitialize = lambda: None
    win32com = types.ModuleType("win32com")
    client = types.ModuleType("win32com.client")
    client.DispatchEx = lambda _name: FakeWord()
    win32com.client = client
    monkeypatch.setitem(sys.modules, "pythoncom", pythoncom)
    monkeypatch.setitem(sys.modules, "win32com", win32com)
    monkeypatch.setitem(sys.modules, "win32com.client", client)

    result = export_word_pdf(source, output)

    # 安全设置降级不再被静默吞掉：结果里逐项列出降级的设置。
    assert result["ok"] is True
    assert result["security_degraded"] == [
        "automation_security",
        "update_links_at_open",
    ]


def test_docx_verifier_turns_render_security_degradation_into_environment_warning(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    _generate_first_revision(workspace, jobs, job)

    class DegradedBroker:
        def run(self, _key, _input, output):
            Path(output).write_bytes(b"fake-pdf")
            return {
                "ok": True,
                "exit_code": 0,
                "stdout": "",
                "stderr": "",
                "security_degraded": ["automation_security"],
            }

    verifier = DocxVerifier(jobs, DegradedBroker())
    _stub_page_rendering(verifier)

    result = verifier.verify(job["id"], version=1, mode="render")

    assert result["passed"] is True
    degradation = [
        item
        for item in result["findings"]
        if item.get("severity") == "warning"
        and item.get("category") == "environment"
        and "automation_security" in item.get("issue", "")
    ]
    assert degradation, result["findings"]
    assert result["render"]["security_degraded"] == ["automation_security"]


def test_vision_qa_sends_page_image_and_returns_structured_findings(tmp_path):
    image = tmp_path / "page-1.png"
    image.write_bytes(b"fake-png")

    class FakeVisionLLM:
        def think(self, messages, **kwargs):
            assert messages[0]["content"][1]["type"] == "image_url"
            return json.dumps(
                {
                    "findings": [
                        {
                            "severity": "blocking",
                            "page": 1,
                            "anchor": "表1",
                            "issue": "表格越界",
                            "suggested_action": "缩小列宽",
                        }
                    ]
                },
                ensure_ascii=False,
            )

    result = VisionQAService(FakeVisionLLM()).inspect_pages([image])

    assert result["available"] is True
    assert result["findings"][0]["severity"] == "blocking"


def test_vision_qa_blocks_invalid_json_instead_of_false_pass(tmp_path):
    image = tmp_path / "page-1.png"
    image.write_bytes(b"fake-png")

    class InvalidVisionLLM:
        def think(self, _messages, **_kwargs):
            return "not-json"

    result = VisionQAService(InvalidVisionLLM()).inspect_pages([image])

    assert result["findings"][0]["severity"] == "blocking"
    assert "findings JSON" in result["findings"][0]["issue"]


def test_vision_qa_retries_invalid_json_once(tmp_path):
    image = tmp_path / "page-1.png"
    image.write_bytes(b"fake-png")

    class RetryVisionLLM:
        def __init__(self):
            self.calls = 0

        def think(self, _messages, **_kwargs):
            self.calls += 1
            return "not-json" if self.calls == 1 else '{"findings":[]}'

    llm = RetryVisionLLM()
    result = VisionQAService(llm).inspect_pages([image])

    assert llm.calls == 2
    assert result == {"available": True, "findings": []}


def test_windows_chinese_font_aliases_are_recognized():
    assert DocxVerifier._font_matches("黑体", {"simhei"})
    assert DocxVerifier._font_matches("宋体", {"simsun"})
    assert DocxVerifier._font_matches("微软雅黑", {"microsoftyahei"})
    assert DocxVerifier._font_matches("楷体_GB2312", {"simkai"})


def test_missing_template_font_is_disclosed_as_warning():
    verifier = DocxVerifier.__new__(DocxVerifier)
    verifier._installed_font_names = lambda: {"simsun"}
    spec = {
        "elements": [
            {
                "runs": [
                    {
                        "effective_font": {
                            "name": "Euclid",
                            "names": {"ascii": "Euclid"},
                        }
                    }
                ]
            }
        ]
    }

    findings = verifier._font_findings(spec)

    assert findings[0]["severity"] == "warning"
    assert "无法声明字体级高保真" in findings[0]["issue"]


def test_content_uses_canonical_outline_title_over_cached_internal_id(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(job["id"], {"tables": [], "unresolved": []})
    content = DocumentContentStore(jobs)
    content.set_outline(
        job["id"],
        {"sections": [{"id": "section-1", "title": "第一章 正确标题"}]},
    )
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "section-1", "section-1", "新的章节正文。")

    sections = content.all_sections(job["id"])

    assert sections[0]["title"] == "第一章 正确标题"


def test_next_version_skips_orphan_revision_directory(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    orphan = jobs.job_root(job["id"]) / "revisions" / "v003"
    orphan.mkdir(parents=True)

    assert jobs.next_version(job["id"]) == 4


def test_partial_verification_does_not_mark_revision_as_fully_qa_passed(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(
        job["id"],
        {"tables": [{"element_id": "body.tbl0000", "action": "preserve"}], "unresolved": []},
    )
    content = DocumentContentStore(jobs)
    content.set_outline(
        job["id"],
        {
            "sections": [
                {"id": "s1", "title": "第一章"},
                {"id": "s2", "title": "第二章"},
            ]
        },
    )
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "s1", "第一章", "正文")
    content.upsert_section(job["id"], "s2", "第二章", "正文")
    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])

    result = DocxVerifier(jobs).verify(job["id"], version=1, mode="deterministic")
    revision = jobs.revision(job["id"], 1)

    assert result["passed"] is True
    assert result["mode"] == "deterministic"
    assert result["qa_report_path"].endswith("qa-report-deterministic.json")
    assert result["render"] is None
    assert result["page_images"] == {"count": 0, "dir": None}
    assert revision["qa_passed"] is False
    assert revision["qa_modes"]["deterministic"]["passed"] is True


def test_full_verification_publishes_only_after_all_qa_passes(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(
        job["id"],
        {"tables": [{"element_id": "body.tbl0000", "action": "preserve"}], "unresolved": []},
    )
    content = DocumentContentStore(jobs)
    content.set_outline(
        job["id"],
        {"sections": [{"id": "s1", "title": "第一章"}, {"id": "s2", "title": "第二章"}]},
    )
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "s1", "第一章", "这是第一章的新正文内容。")
    content.upsert_section(job["id"], "s2", "第二章", "这是第二章的新正文内容。")
    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])
    target = Path(generated["pending_published_path"])
    assert not target.exists()

    class FakeBroker:
        def run(self, _key, _input, output):
            Path(output).write_bytes(b"fake-pdf")
            return {"ok": True, "exit_code": 0, "stdout": "渲" * 5000, "stderr": ""}

    class FakeVision:
        def inspect_pages(self, page_images, template_summary=None, cache_path=None):
            assert page_images
            return {"available": True, "findings": []}

    verifier = DocxVerifier(jobs, FakeBroker(), FakeVision())
    verifier._page_number_findings = lambda _pdf: []

    def fake_render(pdf_path, output_dir):
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page-1.png"
        page.write_bytes(b"fake-png")
        return [page], []

    verifier._render_pdf_pages = fake_render
    result = DocxVerifyTool(verifier).run(
        job["id"], version=1, mode="all"
    )

    assert result["ok"] is True
    assert result["passed"] is True
    assert result["delivery_ready"] is True
    assert Path(result["published_path"]).exists()
    revision = jobs.revision(job["id"], 1)
    assert revision["qa_passed"] is True
    assert revision["delivery_ready"] is True

    # The tool response is a decision-minimal view: no render diagnostics, no
    # per-page image paths, only warning summaries — and it stays small.
    assert set(result.keys()) == {
        "ok",
        "version",
        "passed",
        "delivery_ready",
        "published_path",
        "warning_count",
        "qa_report_path",
        "findings",
    }
    assert all(
        set(item.keys()) == {"anchor", "issue", "suggested_action"}
        for item in result["findings"]
    )
    assert len(json.dumps(result, ensure_ascii=False)) < 4096

    # The persisted report keeps only render tails plus pointers; the complete
    # diagnostics stay in render-result.json on disk.
    report = json.loads(Path(result["qa_report_path"]).read_text(encoding="utf-8"))
    qa_dir = Path(result["qa_report_path"]).parent
    assert report["render"] == {
        "ok": True,
        "returncode": 0,
        "artifacts": [str(qa_dir / "document.pdf")],
        "stdout_tail": "渲" * 500,
        "stderr_tail": "",
        "result_path": str(qa_dir / "render-result.json"),
    }
    assert report["page_images"] == {"count": 1, "dir": str(qa_dir / "pages")}
    full_render = json.loads((qa_dir / "render-result.json").read_text(encoding="utf-8"))
    assert full_render["stdout"] == "渲" * 5000


def test_docx_verify_tool_returns_structured_result_when_all_qa_is_blocked():
    class BlockedVerifier:
        def verify(self, *_args, **_kwargs):
            return {
                "passed": False,
                "version": 3,
                "findings": [
                    {
                        "severity": "blocking",
                        "issue": "目录书签损坏",
                        "anchor": "第一章",
                        "suggested_action": "修复书签后重新验证",
                        "category": "document",
                        "page": 2,
                    },
                    {
                        "severity": "warning",
                        "issue": "模板字体未安装",
                        "anchor": "fonts",
                        "category": "environment",
                    },
                ],
                "render": {"stdout": "x" * 20000, "stderr": "y" * 20000},
                "page_images": {"count": 5, "dir": "qa/v001/pages"},
                "qa_report_path": "qa/v001/qa-report.json",
            }

    result = DocxVerifyTool(BlockedVerifier()).run("job", mode="all")

    assert result["ok"] is False
    assert result["error"] == "DOCX_QA_BLOCKED"
    assert result["version"] == 3
    # Only the blocking [anchor, issue, suggested_action] triples survive; the
    # full report stays on disk at qa_report_path.
    assert result["findings"] == [
        {
            "anchor": "第一章",
            "issue": "目录书签损坏",
            "suggested_action": "修复书签后重新验证",
        }
    ]
    assert result["qa_report_path"] == "qa/v001/qa-report.json"
    assert result["next_action"] == "docx_edit"
    # The verifier report must not be passed through wholesale anymore.
    assert "render" not in result
    assert "page_images" not in result
    assert "blocking_issues" not in result
    assert set(result.keys()) == {
        "ok",
        "error",
        "version",
        "passed",
        "environment_blocked",
        "findings",
        "qa_report_path",
        "next_action",
        "instruction",
    }


def test_docx_verify_tool_flags_environment_blockers():
    class EnvironmentBlockedVerifier:
        def verify(self, *_args, **_kwargs):
            return {
                "passed": False,
                "environment_blocked": True,
                "findings": [
                    {
                        "severity": "blocking",
                        "issue": "视觉模型未配置，未执行逐页图片检查",
                        "category": "environment",
                    }
                ],
            }

    result = DocxVerifyTool(EnvironmentBlockedVerifier()).run("job", mode="all")

    assert result["ok"] is False
    assert result["environment_blocked"] is True
    assert "environment" in result["instruction"]
    # 环境受阻时指引用户同意后的豁免发布路径，而不是继续 docx_edit 死循环。
    assert "waive_environment" in result["instruction"]
    assert result["next_action"] == "docx_verify"
    assert result["findings"] == [
        {
            "anchor": "",
            "issue": "视觉模型未配置，未执行逐页图片检查",
            "suggested_action": "",
        }
    ]


def test_markdown_table_becomes_real_word_table(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(
        job["id"],
        {"tables": [{"element_id": "body.tbl0000", "action": "preserve"}], "unresolved": []},
    )
    content = DocumentContentStore(jobs)
    content.set_outline(
        job["id"],
        {"sections": [{"id": "s1", "title": "第一章"}, {"id": "s2", "title": "第二章"}]},
    )
    jobs.confirm_plan(job["id"])
    content.upsert_section(
        job["id"],
        "s1",
        "第一章",
        "设备清单如下：\n\n| 设备 | 数量 | 负责人 |\n| --- | --- | --- |\n| 基站 | 10 | 李四 |",
    )
    content.upsert_section(job["id"], "s2", "第二章", "正文。")

    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])
    result = Document(generated["docx_path"])
    text = "\n".join(paragraph.text for paragraph in result.paragraphs)

    assert "| --- |" not in text
    assert any(table.rows[1].cells[0].text == "基站" for table in result.tables)


def test_assumption_sections_are_visibly_marked(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(
        job["id"],
        {"tables": [{"element_id": "body.tbl0000", "action": "preserve"}], "unresolved": []},
    )
    content = DocumentContentStore(jobs)
    content.set_outline(
        job["id"],
        {
            "assumptions": ["总投资暂按5000万元测算"],
            "sections": [{"id": "s1", "title": "第一章"}, {"id": "s2", "title": "第二章"}],
        },
    )
    jobs.confirm_plan(job["id"])
    content.upsert_section(
        job["id"], "s1", "第一章", "项目总投资5000万元。", fact_status="assumption"
    )
    content.upsert_section(job["id"], "s2", "第二章", "正文。")

    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])
    text = "\n".join(paragraph.text for paragraph in Document(generated["docx_path"]).paragraphs)

    assert "【假设】" in text


def test_brand_delete_blocks_old_company_left_in_preserved_table(tmp_path):
    template = tmp_path / "brand-table.docx"
    document = Document()
    document.add_paragraph("旧项目报告", style="Title")
    document.add_heading("第一章 旧内容", level=1)
    document.add_paragraph("需要替换的旧正文。")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "项目单位"
    table.rows[0].cells[1].text = "唐山沃盈科技有限公司"
    document.save(template)

    workspace = tmp_path / "workspace"
    session_path = workspace / ".ycore" / "sessions" / "brand"
    session_path.mkdir(parents=True)
    attachments = AttachmentManager(session_path)
    attachment = attachments.import_file(template, role="template")
    jobs = DocumentJobStore(workspace, "brand")
    job = jobs.create(attachment, title="新报告")
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(
        job["id"], {"topic": "新报告", "brand_info": "delete"}, pending_questions=[]
    )
    jobs.set_contract(
        job["id"],
        {"tables": [{"element_id": "body.tbl0000", "action": "preserve"}], "unresolved": []},
    )
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": [{"id": "s1", "title": "第一章 新内容"}]})
    jobs.confirm_plan(job["id"])
    content.upsert_section(job["id"], "s1", "第一章 新内容", "这是新正文。")
    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])

    result = DocxVerifier(jobs).verify(job["id"], version=generated["version"], mode="deterministic")

    assert result["passed"] is False
    assert any("仍残留" in item["issue"] for item in result["findings"])


def _confirm_simple_plan(jobs, job, sections, table_action="preserve"):
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    jobs.set_contract(
        job["id"],
        {"tables": [{"element_id": "body.tbl0000", "action": table_action}], "unresolved": []},
    )
    content = DocumentContentStore(jobs)
    content.set_outline(job["id"], {"sections": sections})
    jobs.confirm_plan(job["id"])
    return content


def test_upsert_section_inherits_confirmed_outline_title(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    content = _confirm_simple_plan(
        jobs,
        job,
        [{"id": "section-1", "title": "第一章 引言"}, {"id": "section-2", "title": "第二章 方法"}],
    )

    missing_title = content.upsert_section(job["id"], "section-1", "", "引言正文。")
    placeholder_title = content.upsert_section(job["id"], "section-2", "section-2", "方法正文。")

    assert missing_title["section"]["title"] == "第一章 引言"
    assert placeholder_title["section"]["title"] == "第二章 方法"


def test_generate_blocks_placeholder_outline_titles(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    # No explicit ids or titles: normalize_outline generates section-N for both.
    content = _confirm_simple_plan(jobs, job, [{"purpose": "第一章"}, {"purpose": "第二章"}])
    content.upsert_section(job["id"], "section-1", "", "第一章正文。")
    content.upsert_section(job["id"], "section-2", "", "第二章正文。")

    with pytest.raises(ValueError, match="PLACEHOLDER_TITLES"):
        DocxBuilder(workspace, jobs, content).generate(job["id"])


def test_generate_rejects_byte_identical_regeneration(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    content = _confirm_simple_plan(
        jobs,
        job,
        [{"id": "s1", "title": "第一章"}, {"id": "s2", "title": "第二章"}],
    )
    content.upsert_section(job["id"], "s1", "第一章", "第一章的新正文。")
    content.upsert_section(job["id"], "s2", "第二章", "第二章的新正文。")
    builder = DocxBuilder(workspace, jobs, content)
    first = builder.generate(job["id"])
    assert first["version"] == 1

    with pytest.raises(ValueError, match="NO_CONTENT_CHANGE"):
        builder.generate(job["id"])

    # The rejected attempt must not leave an orphan revision directory behind.
    revisions_root = jobs.job_root(job["id"]) / "revisions"
    assert sorted(path.name for path in revisions_root.iterdir()) == ["v001"]
    assert jobs.get(job["id"])["current_revision"] == 1

    content.upsert_section(job["id"], "s2", "第二章", "第二章修改后的正文。")
    second = builder.generate(job["id"])
    assert second["version"] == 2


def test_parent_and_optional_sections_render_as_headings(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    content = _confirm_simple_plan(
        jobs,
        job,
        [
            {
                "id": "chapter",
                "title": "第一章 总述",
                "children": [
                    {"id": "leaf-a", "title": "研究背景"},
                    {"id": "leaf-b", "title": "研究意义", "required": False},
                ],
            }
        ],
    )
    # Only the required leaf is written: the parent and the optional leaf stay heading-only.
    content.upsert_section(job["id"], "leaf-a", "研究背景", "研究背景正文。")

    missing = content.get_missing(job["id"])
    assert missing["complete"] is True
    assert "chapter" in missing["present"]
    assert "leaf-b" in missing["present"]

    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])
    texts = [paragraph.text for paragraph in Document(generated["docx_path"]).paragraphs]
    assert "第一章 总述" in texts
    assert "研究意义" in texts

    result = DocxVerifier(jobs).verify(job["id"], version=1, mode="deterministic")
    heading_findings = [
        item for item in result["findings"] if "提纲标题" in str(item.get("issue"))
    ]
    assert heading_findings == []


def test_duplicate_outline_titles_use_expected_occurrence_count(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    content = _confirm_simple_plan(
        jobs,
        job,
        [
            {
                "id": "c1",
                "title": "第一章 现状",
                "children": [{"id": "c1-s", "title": "小结"}],
            },
            {
                "id": "c2",
                "title": "第二章 展望",
                "children": [{"id": "c2-s", "title": "小结"}],
            },
        ],
    )
    content.upsert_section(job["id"], "c1-s", "小结", "第一章小结正文。")
    content.upsert_section(job["id"], "c2-s", "小结", "第二章小结正文。")

    DocxBuilder(workspace, jobs, content).generate(job["id"])
    result = DocxVerifier(jobs).verify(job["id"], version=1, mode="deterministic")

    heading_findings = [
        item for item in result["findings"] if "小结" in str(item.get("anchor"))
    ]
    assert heading_findings == []


def test_set_contract_rejects_unknown_element_ids(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])

    with pytest.raises(ValueError, match="UNKNOWN_CONTRACT_ELEMENT.*body.tbl0099"):
        jobs.set_contract(
            job["id"],
            {"tables": [{"element_id": "body.tbl0099", "action": "rewrite"}]},
        )


def test_reanalysis_is_idempotent_and_preserves_cleared_questions(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    analyzer = DocxTemplateAnalyzer(jobs)
    first = analyzer.analyze(job["id"])
    assert jobs.get(job["id"])["pending_questions"]

    jobs.update_requirements(job["id"], {"topic": "新项目"}, pending_questions=[])
    second = analyzer.analyze(job["id"])

    assert second.get("cached") is True
    assert second["template_spec_path"] == first["template_spec_path"]
    assert jobs.get(job["id"])["pending_questions"] == []


def test_generation_rejects_extra_table_columns(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    content = _confirm_simple_plan(
        jobs,
        job,
        [{"id": "s1", "title": "第一章"}],
        table_action="rewrite",
    )
    content.upsert_section(
        job["id"],
        "s1",
        "第一章",
        "正文。",
        tables=[
            {
                "target_element_id": "body.tbl0000",
                "headers": ["一", "二", "三", "四"],
                "rows": [["1", "2", "3", "4"]],
            }
        ],
    )

    with pytest.raises(ValueError, match="TABLE_COLUMN_MISMATCH"):
        DocxBuilder(workspace, jobs, content).generate(job["id"])


def test_editor_aliases_occurrence_counting_and_strict_targets(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    content = _confirm_simple_plan(
        jobs,
        job,
        [{"id": "s1", "title": "第一章"}, {"id": "s2", "title": "第二章"}],
    )
    content.upsert_section(job["id"], "s1", "第一章", "示例词甲与示例词甲同段出现。")
    content.upsert_section(job["id"], "s2", "第二章", "第二章正文。")
    DocxBuilder(workspace, jobs, content).generate(job["id"])
    editor = DocxEditor(workspace, jobs)

    with pytest.raises(ValueError, match="at least one operation"):
        editor.edit(job["id"], 1, [])

    # Two occurrences inside one paragraph: occurrence counting must see 2, not 1.
    edited = editor.edit(
        job["id"],
        1,
        [
            {
                "operation": "replace_text",
                "old_text": "示例词甲",
                "new_text": "示例词乙",
                "expected_replacements": 2,
            },
            {"operation": "insert_paragraph", "target": "第二章正文。", "content": "追加的说明段。"},
        ],
    )
    assert edited["version"] == 2
    texts = [paragraph.text for paragraph in Document(edited["docx_path"]).paragraphs]
    assert "示例词乙与示例词乙同段出现。" in texts
    assert "追加的说明段。" in texts

    with pytest.raises(ValueError, match="Table target out of range"):
        editor.edit(
            job["id"],
            2,
            [{"operation": "delete", "target": "body.tbl0005"}],
        )
    # Failed edits must not leave orphan revision directories.
    revisions_root = jobs.job_root(job["id"]) / "revisions"
    assert sorted(path.name for path in revisions_root.iterdir()) == ["v001", "v002"]


def test_update_requirements_without_questions_leaves_queue_untouched(document_workspace):
    _workspace, _template, attachments, jobs, job = document_workspace
    DocxTemplateAnalyzer(jobs).analyze(job["id"])
    tool = DocumentJobTool(jobs, attachments)
    questions_before = jobs.get(job["id"])["pending_questions"]
    assert questions_before

    validated = tool.schema.validate(
        {"operation": "update_requirements", "job_id": job["id"], "requirements": {"topic": "新项目"}}
    )
    result = tool.run(**validated)
    assert result["pending_questions"] == questions_before

    validated = tool.schema.validate(
        {
            "operation": "update_requirements",
            "job_id": job["id"],
            "requirements": {},
            "pending_questions": [],
        }
    )
    result = tool.run(**validated)
    assert result["pending_questions"] == []


def test_file_reader_refuses_template_spec(document_workspace, tmp_path):
    workspace, _template, _attachments, jobs, job = document_workspace
    from yc_agents.tools.file_reader import FileReaderTool

    spec_path = jobs.job_root(job["id"]) / "template" / "template-spec.json"
    spec_path.parent.mkdir(parents=True, exist_ok=True)
    spec_path.write_text("{}", encoding="utf-8")

    reader = FileReaderTool(workspace)
    with pytest.raises(PermissionError, match="docx_template_query"):
        reader.run(str(spec_path.relative_to(workspace)), allow_large=True)


def test_update_validates_status_transitions_with_teaching_error(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace

    with pytest.raises(ValueError, match="ILLEGAL_STATUS_TRANSITION") as error:
        jobs.update(job["id"], status="generating")
    message = str(error.value)
    assert "created" in message
    assert "generating" in message
    assert "Allowed" in message

    # Same-status updates pass through untouched.
    assert jobs.update(job["id"], status="created")["status"] == "created"
    with pytest.raises(ValueError, match="Unsupported document job status"):
        jobs.update(job["id"], status="teleporting")
    # force_status stays available as a recovery escape hatch.
    forced = jobs.update(job["id"], status="verifying", force_status=True)
    assert forced["status"] == "verifying"


def test_update_optimistic_concurrency_detects_stale_writers(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    current = jobs.get(job["id"])["updated_at"]

    updated = jobs.update(job["id"], title="并发安全标题", expected_updated_at=current)

    assert updated["title"] == "并发安全标题"
    with pytest.raises(ValueError, match="并发修改"):
        jobs.update(job["id"], title="过期写入", expected_updated_at="2000-01-01T00:00:00")


def test_record_delivery_is_single_source_of_truth_and_append_only(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    content = _confirm_simple_plan(
        jobs, job, [{"id": "s1", "title": "第一章"}, {"id": "s2", "title": "第二章"}]
    )
    content.upsert_section(job["id"], "s1", "第一章", "第一章正文。")
    content.upsert_section(job["id"], "s2", "第二章", "第二章正文。")
    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])
    jobs.update(job["id"], status="waiting_revision")

    with pytest.raises(ValueError, match="Document revision does not exist"):
        jobs.record_delivery(job["id"], 9, published_path="outputs/x.docx")

    delivered = jobs.record_delivery(
        job["id"],
        generated["version"],
        published_path=generated["pending_published_path"],
        qa_report_path="qa/v001/qa-report.json",
        waivers=["用户豁免：页码字体警告"],
    )

    assert delivered["status"] == "delivered"
    raw = json.loads((jobs.job_root(job["id"]) / "job.json").read_text(encoding="utf-8"))
    assert raw["delivery"]["version"] == 1
    assert raw["delivery"]["published_path"] == generated["pending_published_path"]
    assert raw["delivery"]["qa_report_path"] == "qa/v001/qa-report.json"
    assert raw["delivery"]["waivers"] == ["用户豁免：页码字体警告"]
    assert raw["delivery"]["published_at"]
    assert raw["delivery"]["demotions"] == []

    # A delivery is never cleared; failures append demotion events instead.
    with pytest.raises(ValueError, match="DELIVERY_IMMUTABLE"):
        jobs.update(job["id"], delivery=None)
    demoted = jobs.append_delivery_demotion(job["id"], "重验失败：页边距漂移")
    assert demoted["delivery"]["published_path"] == generated["pending_published_path"]
    assert demoted["delivery"]["demotions"][0]["reason"] == "重验失败：页边距漂移"
    assert demoted["delivery"]["demotions"][0]["at"]

    summary = jobs.summary(jobs.get(job["id"]))
    lite = jobs.summary_lite(jobs.get(job["id"]))
    assert summary["delivery"]["demotions"][0]["reason"] == "重验失败：页边距漂移"
    assert summary["demotion_count"] == 1
    assert lite["delivery"] == {
        "version": 1,
        "published_path": generated["pending_published_path"],
        "demotion_count": 1,
    }

    # A delivered job may re-enter the revision loop without losing the record.
    revised = jobs.update(job["id"], status="waiting_revision")
    assert revised["delivery"]["version"] == 1


def test_append_delivery_demotion_requires_existing_delivery(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace

    with pytest.raises(ValueError, match="record_delivery"):
        jobs.append_delivery_demotion(job["id"], "没有交付记录")


def test_job_json_keeps_slim_revision_index_and_manifest_is_authoritative(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    content = _confirm_simple_plan(
        jobs, job, [{"id": "s1", "title": "第一章"}, {"id": "s2", "title": "第二章"}]
    )
    content.upsert_section(job["id"], "s1", "第一章", "第一章正文。")
    content.upsert_section(job["id"], "s2", "第二章", "第二章正文。")
    generated = DocxBuilder(workspace, jobs, content).generate(job["id"])

    raw = json.loads((jobs.job_root(job["id"]) / "job.json").read_text(encoding="utf-8"))
    assert set(raw["revisions"][0]) == {
        "version",
        "docx_sha256",
        "qa_passed",
        "delivery_ready",
        "published_path",
        "manifest_path",
    }
    entry = raw["revisions"][0]
    assert entry["version"] == 1
    assert entry["qa_passed"] is False
    assert entry["published_path"] is None

    manifest = json.loads(Path(entry["manifest_path"]).read_text(encoding="utf-8"))
    assert manifest["docx_path"] == generated["docx_path"]
    assert manifest["package_parts"]
    assert manifest["manifest_path"] == entry["manifest_path"]

    # Store reads hydrate the slim index back from the authoritative manifest.
    hydrated = jobs.revision(job["id"], 1)
    assert hydrated["docx_path"] == generated["docx_path"]
    assert hydrated["package_parts"]
    assert jobs.get(job["id"])["revisions"][0]["docx_path"] == generated["docx_path"]


def test_qa_records_are_stored_as_slim_pointers(document_workspace):
    _workspace, _template, _attachments, jobs, job = document_workspace
    report = {"passed": True, "version": 1, "mode": "all", "findings": [{"issue": "细节" * 200}]}

    jobs.update(job["id"], qa={"v001:all": report, "v001:deterministic": {"passed": False}})

    raw = json.loads((jobs.job_root(job["id"]) / "job.json").read_text(encoding="utf-8"))
    qa_root = jobs.job_root(job["id"]) / "qa"
    assert raw["qa"]["v001:all"] == {
        "passed": True,
        "report_path": str(qa_root / "v001" / "qa-report.json"),
    }
    assert raw["qa"]["v001:deterministic"] == {
        "passed": False,
        "report_path": str(qa_root / "v001" / "qa-report-deterministic.json"),
    }


def _generate_first_revision(workspace, jobs, job):
    content = _confirm_simple_plan(
        jobs, job, [{"id": "s1", "title": "第一章"}, {"id": "s2", "title": "第二章"}]
    )
    content.upsert_section(job["id"], "s1", "第一章", "这是第一章的新正文内容。")
    content.upsert_section(job["id"], "s2", "第二章", "这是第二章的新正文内容。")
    return DocxBuilder(workspace, jobs, content).generate(job["id"])


class _CountingBroker:
    def __init__(self):
        self.calls = 0

    def run(self, _key, _input, output):
        self.calls += 1
        Path(output).write_bytes(b"fake-pdf")
        return {"ok": True, "exit_code": 0, "stdout": "", "stderr": ""}


class _PassVision:
    def __init__(self):
        self.calls = 0

    def inspect_pages(self, page_images, template_summary=None, cache_path=None):
        self.calls += 1
        assert page_images
        return {"available": True, "findings": []}


def _stub_page_rendering(verifier):
    render_calls = []

    def fake_render(pdf_path, output_dir):
        render_calls.append(str(pdf_path))
        output_dir.mkdir(parents=True, exist_ok=True)
        page = output_dir / "page-1.png"
        page.write_bytes(b"fake-png")
        return [page], []

    verifier._render_pdf_pages = fake_render
    verifier._page_number_findings = lambda _pdf: []
    return render_calls


def test_environment_blocked_verify_keeps_status_and_waived_publish_delivers(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    generated = _generate_first_revision(workspace, jobs, job)
    assert jobs.get(job["id"])["status"] == "verifying"

    verifier = DocxVerifier(jobs)  # 无 broker → 环境类阻塞
    result = verifier.verify(job["id"], version=1, mode="all")

    assert result["passed"] is False
    assert result["environment_blocked"] is True
    # 环境故障不再把 job 置为 failed。
    assert jobs.get(job["id"])["status"] == "verifying"

    with pytest.raises(ValueError, match="ENVIRONMENT_BLOCKED"):
        verifier.publish(job["id"], version=1)

    published = verifier.publish(job["id"], version=1, waive_environment=True)

    assert published["published"] is True
    assert published["published_path"] == generated["pending_published_path"]
    assert Path(published["published_path"]).exists()
    assert published["waivers"]
    data = jobs.get(job["id"])
    assert data["status"] == "delivered"
    assert data["delivery"]["version"] == 1
    assert data["delivery"]["published_path"] == published["published_path"]
    assert any("Word" in item or "渲染" in item for item in data["delivery"]["waivers"])
    assert jobs.revision(job["id"], 1)["published_path"] == published["published_path"]

    again = verifier.publish(job["id"], version=1, waive_environment=True)
    assert again["already_published"] is True
    assert again["published_path"] == published["published_path"]
    assert len(jobs.get(job["id"])["delivery"].get("demotions") or []) == 0


def test_publish_requires_full_qa_verdict_and_refuses_document_blockers(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    _generate_first_revision(workspace, jobs, job)
    verifier = DocxVerifier(jobs)

    with pytest.raises(ValueError, match="NO_QA_VERDICT"):
        verifier.publish(job["id"], version=1)

    report_path = jobs.job_root(job["id"]) / "qa" / "v001" / "qa-report.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(
        json.dumps(
            {
                "passed": False,
                "mode": "all",
                "findings": [
                    {
                        "severity": "blocking",
                        "anchor": "第一章",
                        "issue": "目录书签损坏",
                        "category": "document",
                    }
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    jobs.update(job["id"], qa={"v001:all": {"passed": False, "report_path": str(report_path)}})

    with pytest.raises(ValueError, match="PUBLISH_BLOCKED"):
        verifier.publish(job["id"], version=1, waive_environment=True)


def test_full_verify_records_delivery_and_reverify_failure_appends_demotion(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    _generate_first_revision(workspace, jobs, job)
    broker = _CountingBroker()
    verifier = DocxVerifier(jobs, broker, _PassVision())
    _stub_page_rendering(verifier)

    first = verifier.verify(job["id"], version=1, mode="all")

    assert first["passed"] is True
    data = jobs.get(job["id"])
    assert data["status"] == "delivered"
    assert data["delivery"]["version"] == 1
    assert data["delivery"]["published_path"] == first["published_path"]

    class BlockingVision:
        def inspect_pages(self, page_images, template_summary=None, cache_path=None):
            return {
                "available": True,
                "findings": [
                    {
                        "severity": "blocking",
                        "page": 1,
                        "anchor": "表1",
                        "issue": "表格越界",
                        "suggested_action": "缩小列宽",
                    }
                ],
            }

    reverifier = DocxVerifier(jobs, broker, BlockingVision())
    _stub_page_rendering(reverifier)
    second = reverifier.verify(job["id"], version=1, mode="all")

    assert second["passed"] is False
    data = jobs.get(job["id"])
    # 重验失败绝不清空已发布修订的 published_path，改为追加降级事件。
    assert jobs.revision(job["id"], 1)["published_path"] == first["published_path"]
    assert data["delivery"]["published_path"] == first["published_path"]
    assert len(data["delivery"]["demotions"]) == 1
    assert "重验" in data["delivery"]["demotions"][0]["reason"]
    assert Path(first["published_path"]).exists()


def test_verify_write_back_preserves_concurrent_qa_records(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    _generate_first_revision(workspace, jobs, job)

    class SideEffectBroker:
        def run(self, _key, _input, output):
            Path(output).write_bytes(b"fake-pdf")
            data = jobs.get(job["id"])
            jobs.update(
                job["id"],
                qa={
                    **dict(data.get("qa") or {}),
                    "v001:probe": {"passed": True, "report_path": "probe.json"},
                },
            )
            return {"ok": True, "exit_code": 0, "stdout": "", "stderr": ""}

    verifier = DocxVerifier(jobs, SideEffectBroker())
    _stub_page_rendering(verifier)
    result = verifier.verify(job["id"], version=1, mode="render")

    assert result["passed"] is True
    qa = jobs.get(job["id"])["qa"]
    # 渲染窗口期间的并发写不允许被 last-writer-wins 抹掉。
    assert qa["v001:probe"] == {"passed": True, "report_path": "probe.json"}
    assert "v001:render" in qa


def test_render_cache_skips_word_render_for_unchanged_docx(document_workspace):
    workspace, _template, _attachments, jobs, job = document_workspace
    _generate_first_revision(workspace, jobs, job)
    broker = _CountingBroker()
    vision = _PassVision()
    verifier = DocxVerifier(jobs, broker, vision)
    render_calls = _stub_page_rendering(verifier)

    first = verifier.verify(job["id"], version=1, mode="render")

    assert first["passed"] is True
    assert broker.calls == 1
    assert len(render_calls) == 1
    qa_dir = jobs.job_root(job["id"]) / "qa" / "v001"
    manifest = json.loads((qa_dir / "render-cache.json").read_text(encoding="utf-8"))
    assert manifest["docx_sha256"] == jobs.revision(job["id"], 1)["docx_sha256"]
    assert manifest["pdf_path"] == str(qa_dir / "document.pdf")
    assert manifest["page_images"]

    second = verifier.verify(job["id"], version=1, mode="all")

    # 同一版本第二次验证命中 sha256 缓存：不再拉起 Word 渲染子进程和逐页重绘。
    assert broker.calls == 1
    assert len(render_calls) == 1
    assert second["passed"] is True
    assert second["render"]["cache_hit"] is True
    assert second["page_images"]["count"] == 1
    assert vision.calls == 1


def test_vision_page_cache_hits_skip_llm_and_survive_corruption(tmp_path):
    page_one = tmp_path / "page-1.png"
    page_two = tmp_path / "page-2.png"
    page_one.write_bytes(b"png-one")
    page_two.write_bytes(b"png-two")
    cache_path = tmp_path / "vision-cache.json"

    class CountingVisionLLM:
        model = "mimo-vl"

        def __init__(self):
            self.calls = 0

        def think(self, _messages, **_kwargs):
            self.calls += 1
            return json.dumps(
                {"findings": [{"severity": "warning", "anchor": "a", "issue": "贴边"}]},
                ensure_ascii=False,
            )

    llm = CountingVisionLLM()
    service = VisionQAService(llm)

    first = service.inspect_pages([page_one, page_two], cache_path=cache_path)
    assert llm.calls == 2
    assert [item["page"] for item in first["findings"]] == [1, 2]

    second = service.inspect_pages([page_one, page_two], cache_path=cache_path)
    assert llm.calls == 2
    assert second == first

    # 内容换页位后命中同一缓存键，page 编号按当前位置重映射。
    remapped = service.inspect_pages([page_two, page_one], cache_path=cache_path)
    assert llm.calls == 2
    assert [item["page"] for item in remapped["findings"]] == [1, 2]

    cache_path.write_text("{not-json", encoding="utf-8")
    rebuilt = service.inspect_pages([page_one], cache_path=cache_path)
    assert llm.calls == 3
    assert rebuilt["findings"][0]["page"] == 1
    assert json.loads(cache_path.read_text(encoding="utf-8"))["entries"]


def test_vision_environment_failures_are_not_cached(tmp_path):
    page = tmp_path / "page-1.png"
    page.write_bytes(b"png")
    cache_path = tmp_path / "vision-cache.json"

    class InvalidVisionLLM:
        model = "mimo-vl"

        def __init__(self):
            self.calls = 0

        def think(self, _messages, **_kwargs):
            self.calls += 1
            return "not-json"

    llm = InvalidVisionLLM()
    service = VisionQAService(llm)

    first = service.inspect_pages([page], cache_path=cache_path)
    assert first["findings"][0]["category"] == "environment"
    calls_after_first = llm.calls

    second = service.inspect_pages([page], cache_path=cache_path)
    assert second["findings"][0]["category"] == "environment"
    assert llm.calls > calls_after_first


def test_vision_concurrency_is_bounded_and_order_is_stable(tmp_path):
    import threading
    import time

    pages = []
    for index in range(4):
        page = tmp_path / f"page-{index + 1}.png"
        page.write_bytes(f"png-{index}".encode())
        pages.append(page)

    class SlowVisionLLM:
        model = "mimo-vl"

        def __init__(self):
            self.lock = threading.Lock()
            self.active = 0
            self.max_active = 0

        def think(self, messages, **_kwargs):
            with self.lock:
                self.active += 1
                self.max_active = max(self.max_active, self.active)
            time.sleep(0.05)
            with self.lock:
                self.active -= 1
            prompt = messages[0]["content"][0]["text"]
            marker = prompt.split("当前页码：")[1].split("。")[0]
            return json.dumps(
                {"findings": [{"severity": "warning", "anchor": f"p{marker}", "issue": "x"}]},
                ensure_ascii=False,
            )

    llm = SlowVisionLLM()
    result = VisionQAService(llm, max_workers=2).inspect_pages(pages)

    assert llm.max_active <= 2
    assert [item["page"] for item in result["findings"]] == [1, 2, 3, 4]
    assert [item["anchor"] for item in result["findings"]] == ["p1", "p2", "p3", "p4"]


def test_docx_verify_tool_publish_operation_and_validation():
    class PublishVerifier:
        def __init__(self):
            self.calls = []

        def publish(self, job_id, version=None, waive_environment=False):
            self.calls.append((job_id, version, waive_environment))
            return {
                "published": True,
                "already_published": False,
                "version": 2,
                "published_path": "outputs/x.docx",
                "waivers": ["视觉模型未配置，未执行逐页图片检查"],
            }

    verifier = PublishVerifier()
    result = DocxVerifyTool(verifier).run(
        "job", version=2, operation="publish", waive_environment=True
    )

    assert verifier.calls == [("job", 2, True)]
    assert result["ok"] is True
    assert result["operation"] == "publish"
    assert result["published_path"] == "outputs/x.docx"
    assert result["waivers"] == ["视觉模型未配置，未执行逐页图片检查"]
    assert "豁免" in result["instruction"]

    with pytest.raises(ValueError, match="operation"):
        DocxVerifyTool(verifier).run("job", operation="teleport")
